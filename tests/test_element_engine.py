# -*- coding: utf-8 -*-
"""The element identification engine (elementid) and the fixes made with it.

  * group A - doors: the indicators, the key combinations; a window or a
    cut-out is Step 3's (the X sequence), recorded as drawn, never scored;
  * group B - staircases (tier 1) and lift lobby, service lobby, lift
    machine room (tier 2, gated on tier 1); a shaft, a lift or a void is
    Step 3's, a fixed record (2026-10-16);
  * bands, actions, tie-breaks, the records and the workbook;
  * the drawing: an X in a dashed outline in a wall is a cut-out, the
    elements are tagged and the doubtful ones marked on AI_REVIEW;
  * the doors (a leaf shut in a wall band, a lift's entrance), the cut-out
    in a wall run, the shaft's room to its door, the balcony's end walls,
    a wall's end never carried across a doorway (FHC-3), LIFT PRE a shaft.

Run:  python -m unittest discover -s tests -v
"""

import math
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import ezdxf  # noqa: E402
from shapely.geometry import Polygon, box  # noqa: E402
from shapely.ops import unary_union  # noqa: E402

from Drawing_Analyzer import dxf_reader as dr  # noqa: E402
from Drawing_Analyzer import elementid as E  # noqa: E402
from Drawing_Analyzer import elements as el  # noqa: E402
from Drawing_Analyzer import layernames as LN  # noqa: E402
from Drawing_Analyzer import rules  # noqa: E402


def seg(x1, y1, x2, y2, layer="A-WALL", door=False):
    return (x1, y1, x2, y2, layer, math.hypot(x2 - x1, y2 - y1), -1, door, False)


def rect(x0, y0, x1, y1):
    return [(x0, y0), (x1, y0), (x1, y1), (x0, y1)]


def outline(x0, y0, x1, y1, layer):
    p = rect(x0, y0, x1, y1)
    return [seg(p[k][0], p[k][1], p[(k + 1) % 4][0], p[(k + 1) % 4][1], layer) for k in range(4)]


def x_mark(x0, y0, x1, y1, layer):
    return [seg(x0, y0, x1, y1, layer), seg(x0, y1, x1, y0, layer)]


def keys(segs):
    return frozenset(dr._seg_key(s) for s in segs)


class BandTests(unittest.TestCase):

    def test_bands_and_actions(self):
        self.assertEqual(E.band(95), "High")
        self.assertEqual(E.band(80), "High")
        self.assertEqual(E.band(79), "Medium")
        self.assertEqual(E.band(50), "Medium")
        self.assertEqual(E.band(49), "Low")
        self.assertEqual(E.band(20), "Low")
        self.assertEqual(E.band(19), "Unresolved")
        self.assertEqual(E.action(50), "Auto-classify")
        self.assertEqual(E.action(35), "Needs review")
        self.assertEqual(E.action(10), "Hold - unresolved")
        self.assertEqual(E.action(90, E.UNRESOLVED), "Hold - unresolved")

    def test_layer_kinds(self):
        self.assertEqual(E.layer_kind("A-DOOR"), "door")
        self.assertEqual(E.layer_kind("XREF$0$A-GLAZING"), "glass")
        self.assertEqual(E.layer_kind("A3-CUTOUTS"), "cutout")
        self.assertEqual(E.layer_kind("A-STRS"), "stair")
        self.assertEqual(E.layer_kind("A-ELEV 2"), "lift")
        self.assertEqual(E.layer_kind("A-WALL-MAS"), "wall")
        self.assertEqual(E.layer_kind("A-FURN"), "other")


class OutlineTests(unittest.TestCase):

    def test_dashed_and_solid_outlines(self):
        cut = box(1000, 0, 1600, 200)
        dashed = outline(1000, 0, 1600, 200, "A-CUTOUT") + x_mark(1000, 0, 1600, 200, "A-CUTOUT")
        ev = E.Evidence(dashed, keys(dashed))
        self.assertEqual(E.boundary_type(cut, ev)[0], "dashed")
        ev = E.Evidence(dashed)
        self.assertEqual(E.boundary_type(cut, ev)[0], "solid")

    def test_wall_faces_and_area_overlays_are_not_the_outline(self):
        """A wall's faces running past the opening and an area take-off
        polyline are not its outline: its own dashed X decides."""
        cut = box(1000, 0, 1600, 200)
        segs = [seg(0, 0, 5000, 0, "A-WALL"), seg(0, 200, 5000, 200, "A-WALL")] + \
            outline(1000, 0, 1600, 200, "CT-AREA")
        xm = x_mark(1000, 0, 1600, 200, "A3-CUTOUTS")
        ev = E.Evidence(segs + xm, keys(xm))
        self.assertEqual(E.boundary_type(cut, ev)[0], "dashed")

    def test_x_found(self):
        cut = box(0, 0, 900, 300)
        ev = E.Evidence(x_mark(0, 0, 900, 300, "A-CUTOUT"))
        self.assertTrue(E.has_x(cut, ev))
        self.assertFalse(E.has_x(cut, E.Evidence([seg(0, 0, 900, 300, "A-CUTOUT")])))


def wall_with_gap(x0, x1, y0=0.0, t=200.0, run=(-3000.0, 6000.0)):
    """A wall along x with an opening x0..x1."""
    return unary_union([box(run[0], y0, x0, y0 + t), box(x1, y0, run[1], y0 + t)])


class GroupATests(unittest.TestCase):

    def facts(self, o, segs, dashed=(), arcs=(), texts=(), walls=None):
        ev = E.Evidence(segs, dashed, arcs, texts)
        f = E._opening_facts(o, ev, walls)
        return f, {E.DOOR: E.score_door(f)}

    def test_a_framed_hinged_door_is_high(self):
        walls = wall_with_gap(1000, 1900)
        quad = box(1000, 0, 1900, 200)
        segs = [seg(1000, 0, 1000, 200, "A-WALL"), seg(1900, 0, 1900, 200, "A-WALL"),     # jamb lines
                seg(1000, 20, 1000, 180, "A-DOOR", True), seg(1900, 20, 1900, 180, "A-DOOR", True),
                seg(1000, 200, 1000, 1100, "A-DOOR", True)]                              # the leaf, open
        arcs = [(1000.0, 200.0, 900.0)]
        f, sc = self.facts({"poly": quad, "kind": "door"}, segs, arcs=arcs, texts=[{"text": "D1", "x": 1450, "y": 600}],
                           walls=walls)
        winner, score, tie, _n = E.decide_a(f, sc)
        self.assertEqual(winner, E.DOOR)
        self.assertGreaterEqual(score, 90)
        self.assertEqual(list(sc), [E.DOOR])                        # only the door is scored

    def test_frame_leaf_opening_without_swing_is_medium(self):
        walls = wall_with_gap(1000, 1900)
        segs = [seg(1000, 20, 1000, 180, "A-DOOR", True), seg(1900, 20, 1900, 180, "A-DOOR", True),
                seg(1000, 90, 1900, 90, "A-DOOR", True)]
        f, sc = self.facts({"poly": box(1000, 0, 1900, 200), "kind": "door"}, segs, walls=walls)
        self.assertEqual(int(sc[E.DOOR].score), 75)

    def test_swing_alone_off_the_wall_is_rejected(self):
        segs = []
        f, sc = self.facts({"poly": box(1000, 0, 1900, 200), "kind": "door"}, segs, arcs=[(1000.0, 200.0, 900.0)],
                           walls=None)
        self.assertTrue(sc[E.DOOR].negative)
        self.assertEqual(sc[E.DOOR].score, 0)

    def test_a_lift_entrance_slides(self):
        walls = wall_with_gap(700, 1800)
        quad = box(700, 0, 1800, 200)
        segs = [seg(700, 0, 700, 200, "A-SLIP FORM RC"), seg(1800, 0, 1800, 200, "A-SLIP FORM RC"),
                seg(700, 0, 1800, 0, "A-ELEV 2"), seg(700, 200, 1800, 200, "A-ELEV 2")]
        f, sc = self.facts({"poly": quad, "kind": "door", "door_kind": "lift"}, segs, walls=walls)
        self.assertTrue(f["sliding"])
        self.assertGreaterEqual(sc[E.DOOR].score, 85)

    def test_windows_and_cut_outs_are_step3s_fixed(self):
        """2026-10-16: a window or a cut-out is decided in Step 3 (the X
        sequence's steps 4 and 5); the engine records it as drawn - score
        100, never reviewed - with the step and the reason from its tag."""
        walls = wall_with_gap(1000, 2200)
        cut, win = box(1000, 0, 2200, 200), box(3000, 0, 4200, 200)
        segs = outline(1000, 0, 2200, 200, "A-MARK") + x_mark(1000, 0, 2200, 200, "A-MARK")
        ops = [{"poly": cut, "kind": "cutout", "x": True, "current_class": E.CUTOUT,
                "x_tag": ("cutout", 5, "inside a wall, no glazing, window frame or window tag")},
               {"poly": win, "kind": "window", "current_class": E.WINDOW}]
        recs = E.identify_floor([], ops, walls, E.Evidence(segs), floor="F", prefix="F1")
        by = {r["classification"]: r for r in recs}
        self.assertEqual(set(by), {E.CUTOUT, E.WINDOW})
        for r in recs:
            self.assertEqual(r["confidence_score"], 100)
            self.assertEqual(r["action"], E.FIXED_ACTION)
            self.assertTrue(r["fixed"])
            self.assertFalse(r["review"])
        self.assertIn("step 5", by[E.CUTOUT]["notes"])
        self.assertIn("no glazing", by[E.CUTOUT]["notes"])
        self.assertIn("Step 3", by[E.WINDOW]["notes"])

    def test_glazing_frame_and_panes_are_read(self):
        walls = wall_with_gap(1000, 2800)
        quad = box(1000, 0, 2800, 200)
        segs = [seg(1000, 0, 1000, 200, "A-WINDOW"), seg(2800, 0, 2800, 200, "A-WINDOW"),
                seg(1000, 60, 2800, 60, "A-GLAZING"), seg(1000, 100, 2800, 100, "A-GLAZING"),
                seg(1000, 140, 2800, 140, "A-GLAZING")]
        f, _sc = self.facts({"poly": quad, "kind": "window"}, segs, walls=walls)
        self.assertTrue(f["glazing"])
        self.assertTrue(f["panes"])
        self.assertTrue(f["in_wall"])

    def test_a_slab_opening_beside_a_shaft_door_has_no_leaf_of_its_own(self):
        """The door of the shaft beside the X strip is not the strip's."""
        x = box(0, 0, 1200, 300)
        segs = outline(0, 0, 1200, 300, "A-CUTOUT") + x_mark(0, 0, 1200, 300, "A-CUTOUT") + [
            seg(100, 320, 100, 400, "A-DOOR", True), seg(1000, 320, 1000, 400, "A-DOOR", True),
            seg(100, 360, 1000, 360, "A-DOOR", True)]
        f, _sc = self.facts({"poly": x, "kind": "x", "x": True, "name": "PHE SHAFT"}, segs, walls=None)
        self.assertFalse(f["leaf"])
        self.assertEqual(f["x_context"], "shaft")

    def test_a_weak_door_is_unresolved(self):
        f = {"x": False, "boundary": "unclear", "in_wall": False}
        a = E._Score(E.DOOR)
        a.score = 10.0
        winner, score, tie, notes = E.decide_a(f, {E.DOOR: a})
        self.assertEqual(winner, E.UNRESOLVED)
        a.score = 35.0
        self.assertEqual(E.decide_a(f, {E.DOOR: a})[0], E.DOOR)    # Low: a possible door, for review


def treads(x0, x1, y0, n, step=300.0, layer="A-STRS"):
    return [seg(x0, y0 + k * step, x1, y0 + k * step, layer) for k in range(n)]


class GroupBTests(unittest.TestCase):

    def run_floor(self, spaces, openings=(), segs=(), texts=(), walls=None, dashed=()):
        ev = E.Evidence(list(segs), dashed, (), list(texts))
        recs = E.identify_floor(spaces, list(openings), walls, ev, floor="F", prefix="F1")
        return {r["name"]: r for r in recs if r["group"] == "B"}, recs

    def test_a_shaft_drawn_in_step3_is_recorded_as_drawn(self):
        sp = [{"poly": box(0, 0, 600, 600), "name": "TEN", "family": "MEP / Services", "current_class": E.SHAFT},
              {"poly": box(800, 0, 1400, 600), "name": "FAPA", "family": "MEP / Services", "current_class": E.SHAFT}]
        got, _ = self.run_floor(sp)
        self.assertEqual(got["TEN"]["classification"], E.SHAFT)
        self.assertEqual(got["TEN"]["confidence_score"], 100)
        self.assertTrue(got["TEN"]["agrees"])
        self.assertTrue(got["TEN"]["fixed"])

    def test_an_x_with_nothing_round_it_is_the_x_sequences(self):
        """2026-10-16: no X is scored again - an X shaft, lift or void is
        what the X sequence made it, even with nothing round it."""
        sp = [{"poly": box(0, 0, 600, 600), "name": "Unknown Shaft", "family": "Shaft", "current_class": E.SHAFT,
               "x": True, "x_tag": ("shaft", 2, "a wall within 0 mm of its S side")}]
        got, _ = self.run_floor(sp, segs=x_mark(0, 0, 600, 600, "A-CUTOUT"))
        r = got["Unknown Shaft"]
        self.assertEqual(r["classification"], E.SHAFT)
        self.assertEqual(r["action"], E.FIXED_ACTION)
        self.assertFalse(r["review"])
        self.assertIn("step 2", r["notes"])

    def test_a_lift_with_its_door_and_lobby_then_the_lobby(self):
        lift = box(0, 0, 2500, 2500)
        lobby = box(-1000, -3200, 6000, -200)
        door = box(700, -200, 1800, 0)
        sp = [{"poly": lift, "name": "PASSENGER LIFT 13 PAX", "family": "Lift", "current_class": E.LIFT},
              {"poly": lobby, "name": "LIFT LOBBY", "family": "Lift Lobby", "current_class": E.LIFT_LOBBY}]
        walls = unary_union([box(-200, -200, 2700, 2700).difference(lift).difference(door)])
        op = [{"poly": door, "kind": "door", "door_kind": "lift", "current_class": E.DOOR}]
        got, _ = self.run_floor(sp, op, walls=walls)
        self.assertEqual(got["PASSENGER LIFT 13 PAX"]["classification"], E.LIFT)
        self.assertGreaterEqual(got["PASSENGER LIFT 13 PAX"]["confidence_score"], 90)
        lob = got["LIFT LOBBY"]
        self.assertEqual(lob["classification"], E.LIFT_LOBBY)
        self.assertEqual(lob["tier"], "2")
        self.assertTrue(lob["prerequisite_met"])
        self.assertTrue(any(p["class"] == E.LIFT for p in lob["prerequisite_elements"]))

    def test_no_lift_no_lift_lobby(self):
        sp = [{"poly": box(0, 0, 4000, 3000), "name": "LIFT LOBBY", "family": "Lift Lobby",
               "current_class": E.LIFT_LOBBY}]
        got, _ = self.run_floor(sp)
        r = got["LIFT LOBBY"]
        self.assertNotEqual(r["classification"], E.LIFT_LOBBY)
        self.assertFalse(r["prerequisite_met"])

    def test_a_staircase_and_the_void_between_its_flights(self):
        stair = box(0, 0, 4500, 7000)
        segs = treads(0, 2000, 1500, 12) + treads(2500, 4500, 1500, 12)
        void = box(2000, 1500, 2500, 4800)
        sp = [{"poly": stair, "name": "ST-01", "family": "Staircase", "current_class": E.STAIR},
              {"poly": void, "name": "VOID", "family": "Void", "current_class": E.VOID, "x": True,
               "x_tag": ("void", 7, "nothing above claimed it: a stair's well")}]
        got, _ = self.run_floor(sp, segs=segs + x_mark(2000, 1500, 2500, 4800, "A-CUTOUT"),
                                texts=[{"text": "UP", "x": 1000, "y": 3000}])
        st = got["ST-01"]
        self.assertEqual(st["classification"], E.STAIR)
        self.assertGreaterEqual(st["confidence_score"], 90)
        v = got["VOID"]
        self.assertEqual(v["classification"], E.VOID)          # the X sequence's step 7
        self.assertTrue(v["fixed"])
        self.assertFalse(v["review"])
        self.assertIn("stair's well", v["notes"])

    def test_flights_found_side_by_side(self):
        stair = box(0, 0, 4500, 7000)
        ev = E.Evidence(treads(0, 2000, 1500, 12) + treads(2500, 4500, 1500, 12)
                        + [seg(0, 1200, 4500, 1200, "A-STRS")])          # a landing edge across both
        fl = E.flights_in(stair, ev)
        self.assertEqual(sorted(f["n"] for f in fl), [12, 12])

    def test_lift_pre_is_a_shaft(self):
        sp = [{"poly": box(0, 0, 1200, 1500), "name": "LIFT PRE", "family": "Shaft", "current_class": E.SHAFT,
               "x": True},
              {"poly": box(1400, 0, 3900, 2500), "name": "24 PAX P.LIFT", "family": "Lift", "current_class": E.LIFT},
              {"poly": box(0, -900, 1200, -200), "name": "FHC", "family": "Fire / Life Safety",
               "current_class": E.SHAFT}]
        got, _ = self.run_floor(sp, segs=x_mark(0, 0, 1200, 1500, "A-CUTOUT"))
        self.assertEqual(got["LIFT PRE"]["classification"], E.SHAFT)
        self.assertGreaterEqual(got["LIFT PRE"]["confidence_score"], 75)

    def test_a_machine_room_with_a_door_is_not_one(self):
        lift = box(0, 0, 2500, 2500)
        mr = box(2700, 0, 5000, 2500)
        sp = [{"poly": lift, "name": "LIFT", "family": "Lift", "current_class": E.LIFT},
              {"poly": mr, "name": "LIFT MACHINE ROOM", "family": "Machine Room", "current_class": E.MACHINE_ROOM}]
        op = [{"poly": box(3500, 2500, 4400, 2700), "kind": "door", "door_kind": "exit", "current_class": E.DOOR}]
        got, _ = self.run_floor(sp, op)
        r = got["LIFT MACHINE ROOM"]
        self.assertNotEqual(r["classification"], E.MACHINE_ROOM)
        self.assertTrue(r["negative_rule_triggered"])

    def test_records_carry_the_specified_fields(self):
        sp = [{"poly": box(0, 0, 600, 600), "name": "TEN", "family": "MEP / Services", "current_class": E.SHAFT}]
        _got, recs = self.run_floor(sp)
        r = E.public(recs[0])
        for k in ("element_id", "group", "tier", "classification", "confidence_score", "confidence_band",
                  "indicators_matched", "indicators_missing", "boundary_type", "x_symbol_present",
                  "x_symbol_context", "prerequisite_elements", "prerequisite_met", "negative_rule_triggered",
                  "tie_break_applied", "competing_classes", "action", "notes"):
            self.assertIn(k, r)
        self.assertFalse(any(k.startswith("_") for k in r))
        self.assertTrue(all({"name", "weight"} <= set(m) for m in r["indicators_matched"]))

    def test_the_workbook(self):
        from openpyxl import load_workbook
        sp = [{"poly": box(0, 0, 600, 600), "name": "TEN", "family": "MEP / Services", "current_class": E.SHAFT}]
        _got, recs = self.run_floor(sp)
        path = os.path.join(tempfile.mkdtemp(), "e.xlsx")
        E.write_elements_xlsx([E.public(r) for r in recs], path)
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Elements", "Summary"])
        self.assertEqual(wb["Elements"].max_row, 1 + len(recs))


class DrawingTests(unittest.TestCase):
    """A room with two X-marked openings in its walls - one outlined
    dashed, one solid - and a door, through annotate_dxf."""

    @classmethod
    def setUpClass(cls):
        tmp = tempfile.mkdtemp()
        doc = ezdxf.new()
        doc.header["$INSUNITS"] = 4
        if "HIDDEN" not in doc.linetypes:
            doc.linetypes.add("HIDDEN", pattern=[0.375, 0.25, -0.125])
        msp = doc.modelspace()
        plate = rect(0, 0, 40000, 30000)
        msp.add_lwpolyline(plate, close=True, dxfattribs={"layer": "A-FLOOR"})

        def r(x0, y0, x1, y1, layer="A-WALL", lt=None):
            at = {"layer": layer}
            if lt:
                at["linetype"] = lt
            msp.add_lwpolyline(rect(x0, y0, x1, y1), close=True, dxfattribs=at)
        r(9800, 9800, 11000, 10000)
        r(12200, 9800, 14000, 10000)
        r(15200, 9800, 18200, 10000)
        r(9800, 16000, 18200, 16200)
        r(9800, 10000, 10000, 16000)
        r(18000, 10000, 18200, 12000)
        r(18000, 12900, 18200, 16000)
        for x0, x1, lt in ((11000, 12200, "HIDDEN"), (14000, 15200, None)):
            r(x0, 9800, x1, 10000, "A-WINDOW", lt)
            at = {"layer": "A-WINDOW"}
            if lt:
                at["linetype"] = lt
            msp.add_line((x0, 9800), (x1, 10000), dxfattribs=at)
            msp.add_line((x0, 10000), (x1, 9800), dxfattribs=at)
        msp.add_arc((18000, 12000), 900, 90, 180, dxfattribs={"layer": "A-DOOR"})
        msp.add_line((18000, 12000), (18000, 12900), dxfattribs={"layer": "A-DOOR"})
        msp.add_mtext("OFFICE", dxfattribs={"insert": (14000, 13000), "char_height": 200, "layer": "A-TEXT"})
        src = os.path.join(tmp, "in.dxf")
        doc.saveas(src)
        floors = [{"name": "F1", "polygon": plate}]
        rows = dr.extract_texts(src, floors)
        cls.out = os.path.join(tmp, "out.dxf")
        cls.stats = dr.annotate_dxf(src, cls.out, floors, rows)
        cls.doc = ezdxf.readfile(cls.out)
        cls.msp = cls.doc.modelspace()
        cls.areas = dr.compute_room_areas(cls.out, "AI_FLOOR")

    def polys(self, layer):
        out = []
        for e in self.msp.query('LWPOLYLINE[layer=="%s"]' % layer):
            out.append(Polygon([(p[0], p[1]) for p in e.get_points("xy")]).buffer(0))
        return out

    def test_the_dashed_x_is_a_cut_out_the_solid_one_a_window(self):
        cuts, wins = self.polys(LN.CUTOUT), self.polys(LN.WINDOW)
        dashed, solid = box(11000, 9800, 12200, 10000), box(14000, 9800, 15200, 10000)
        self.assertTrue(any(c.intersection(dashed).area > 0.8 * dashed.area for c in cuts))
        self.assertFalse(any(w.intersection(dashed).area > 0.2 * dashed.area for w in wins))
        self.assertTrue(any(w.intersection(solid).area > 0.8 * solid.area for w in wins))
        self.assertEqual(self.stats.get("windows_dashed_to_cutouts"), 1)

    def test_every_element_scored_and_tagged(self):
        recs = self.stats.get("elements")
        self.assertTrue(recs)
        classes = {r["classification"] for r in recs if r["group"] == "A"}
        self.assertIn(E.CUTOUT, classes)
        self.assertIn(E.WINDOW, classes)
        self.assertIn(E.DOOR, classes)
        tagged = [e for e in self.msp if e.has_xdata(E.ELEMENT_APPID)]
        self.assertGreaterEqual(len(tagged), 3)
        self.assertTrue(all(len(t) == 5 for e in tagged for t in E.element_tag(e)))
        self.assertIn("elements_summary", self.stats)

    def test_review_marks_are_never_rooms_or_labels(self):
        self.assertTrue(LN.is_reserved(LN.REVIEW))
        self.assertFalse(LN.is_room_layer(LN.REVIEW))
        names = {r["name"] for r in self.areas["rooms"]}
        self.assertFalse(any(n and n.startswith("F1-") for n in names))

    def test_the_rules_exist(self):
        d = rules.defaults()
        self.assertNotIn("dashed_x_cutouts", d)            # the X sequence's, not a switch (2026-10-16)
        self.assertTrue(d["element_scores"])


class ReviewTextTests(unittest.TestCase):

    def test_notes_on_ai_review_are_not_read(self):
        doc = ezdxf.new()
        msp = doc.modelspace()
        msp.add_mtext("F1-B001 Void 95", dxfattribs={"layer": LN.REVIEW, "insert": (0, 0)})
        msp.add_mtext("OFFICE", dxfattribs={"layer": "A-TEXT", "insert": (0, 0)})
        got = [t.text for t in dr._collect_texts_in_modelspace(msp)]
        self.assertEqual(got, ["OFFICE"])


class StackedLabelTests(unittest.TestCase):

    def test_two_rooms_labels_are_not_one(self):
        """A shaft's cut-out and its room to the door carry the same name,
        one above the other: two labels, not "PHE SHAFT PHE SHAFT"."""
        T = dr.TextRecord
        recs = [T(text="PHE SHAFT", source="MTEXT", layer="AI_SHAFT", x=0.0, y=0.0, rotation=0.0, height=150.0,
                  handles=("a",)),
                T(text="PHE SHAFT", source="MTEXT", layer="AI_SHAFT", x=0.0, y=-250.0, rotation=0.0, height=150.0,
                  handles=("b",))]
        got = sorted(r.text for r in dr._merge_stacked_lines(recs))
        self.assertEqual(got, ["PHE SHAFT", "PHE SHAFT"])

    def test_a_drawing_label_in_two_lines_is_still_one(self):
        T = dr.TextRecord
        recs = [T(text="24 PAX", source="TEXT", layer="A-TEXT", x=0.0, y=0.0, rotation=0.0, height=150.0, handles=("a",)),
                T(text="P.LIFT", source="TEXT", layer="A-TEXT", x=0.0, y=-250.0, rotation=0.0, height=150.0,
                  handles=("b",))]
        self.assertEqual([r.text for r in dr._merge_stacked_lines(recs)], ["24 PAX P.LIFT"])


class DashedLinesTests(unittest.TestCase):

    def test_dashed_segments_are_known(self):
        doc = ezdxf.new()
        if "HIDDEN" not in doc.linetypes:
            doc.linetypes.add("HIDDEN", pattern=[0.375, 0.25, -0.125])
        doc.layers.add("CUT", linetype="HIDDEN")
        msp = doc.modelspace()
        msp.add_line((0, 0), (1000, 0), dxfattribs={"layer": "CUT"})
        msp.add_line((0, 500), (1000, 500), dxfattribs={"layer": "A-WALL"})
        msp.add_line((0, 900), (1000, 900), dxfattribs={"layer": "A-WALL", "linetype": "HIDDEN"})
        self.assertIn("HIDDEN", dr._dashed_linetypes(doc))
        self.assertNotIn("CONTINUOUS", dr._dashed_linetypes(doc))
        segs = dr._collect_raw_segs(doc, msp)
        d = {s[1]: dr._seg_key(s) in segs.dashed for s in segs}
        self.assertEqual(d, {0.0: True, 500.0: False, 900.0: True})


class DoorBridgeTests(unittest.TestCase):

    def test_a_wall_end_is_not_carried_across_the_room(self):
        """FHC-3 on P285: the end of the wall beside a door ran on across the
        room behind it; only the faces along the wall cross the doorway."""
        faces = [(0, 0, 900, 0), (0, 250, 900, 250), (900, 0, 900, 250),        # the left wall and its end
                 (1800, 0, 3000, 0), (1800, 250, 3000, 250), (1800, 0, 1800, 250),
                 (800, -550, 1900, -550)]                                       # the cut-out's edge behind
        swings = [(900.0, 0.0, 900.0), (1000.0, 120.0, 560.0)]
        got = dr._door_bridges(faces, [], swings)
        horiz = [b for b in got if abs(b[1] - b[3]) < 1]
        vert = [b for b in got if abs(b[0] - b[2]) < 1]
        self.assertTrue(any(abs(min(b[0], b[2]) - 900) < 1 and abs(max(b[0], b[2]) - 1800) < 1 for b in horiz))
        self.assertFalse(any(abs(b[0] - 900) < 1 for b in vert))


    def test_a_door_between_two_wall_ends_still_closes(self):
        """H10: shaft doors between the ends of the partitions: the ends,
        carried on, meet the next partition at its corner."""
        faces = [(0, 0, 0, 1175), (200, 0, 200, 1175), (0, 0, 200, 0),
                 (1400, 0, 1400, 1175), (1600, 0, 1600, 1175), (1400, 0, 1600, 0)]
        got = dr._door_bridges(faces, [], [(200.0, 0.0, 1200.0)])
        self.assertTrue(any(abs(b[1]) < 1 and abs(b[3]) < 1 and abs(min(b[0], b[2]) - 200) < 1
                            and abs(max(b[0], b[2]) - 1400) < 1 for b in got))


class LeafDoorTests(unittest.TestCase):

    def test_a_leaf_shut_in_a_wall_band_opens_it(self):
        """P285 FD4: the masonry wall runs on through the door; its leaf lies
        shut in the band, framed at both ends."""
        strips = [{"pts": rect(0, 0, 3000, 200), "t": 200, "layers": ("A-WALL", "A-WALL")}]
        segs = [seg(1000, 90, 1800, 90, "A-DOOR", True), seg(1000, 130, 1800, 130, "A-DOOR", True),
                seg(990, 40, 990, 180, "A-DOOR", True), seg(1810, 40, 1810, 180, "A-DOOR", True)]
        quads = dr._leaf_door_quads(segs, strips)
        self.assertEqual(len(quads), 1)
        g = Polygon(quads[0])
        self.assertLess(abs(g.bounds[1] - 0), 1)
        self.assertLess(abs(g.bounds[3] - 200), 1)
        self.assertTrue(g.contains(box(1000, 10, 1800, 190)))

    def test_no_frame_no_door(self):
        strips = [{"pts": rect(0, 0, 3000, 200), "t": 200, "layers": ("A-WALL", "A-WALL")}]
        segs = [seg(1000, 90, 1800, 90, "A-DOOR", True), seg(1000, 130, 1800, 130, "A-DOOR", True)]
        self.assertEqual(dr._leaf_door_quads(segs, strips), [])


class LiftDoorTests(unittest.TestCase):

    def test_the_gap_in_the_lift_wall_is_its_door(self):
        lift = box(0, 0, 2500, 2500)
        gap = box(700, -200, 1800, 0)
        walls = box(-200, -200, 2700, 2700).difference(lift).difference(gap)
        got = el.lift_door_quads([lift], walls)
        self.assertEqual(len(got), 1)
        q = Polygon(got[0][0])
        self.assertGreater(q.intersection(gap).area, 0.8 * gap.area)

    def test_lift_doors_are_primary(self):
        """2026-10-17: a lift's door stands between two spaces (the lift and
        the floor in front of it): primary, on AI_P_DOORS."""
        lev = el.door_levels([[(700, -200), (1800, -200), (1800, 0), (700, 0)]],
                             [box(0, 0, 2500, 2500), box(-1000, -3000, 4000, -200)],
                             [False, True], [False, False], kinds=[(el.LIFT, 0, 1)])
        self.assertEqual(lev[0][0], "primary")

    def test_a_lift_open_its_whole_front_has_that_front_as_its_door(self):
        """Rule lift_without_door_mechanical: no gap in a wall, but the
        whole front open, the other three sides walled - the front is the
        door (only with open_front)."""
        lift = box(0, 0, 2500, 2500)
        walls = box(-200, 0, 2700, 2700).difference(lift)          # nothing along y = 0
        self.assertEqual(el.lift_door_quads([lift], walls), [])
        got = el.lift_door_quads([lift], walls, open_front=True)
        self.assertEqual(len(got), 1)
        q = Polygon(got[0][0])
        self.assertTrue(got[0][1].get("open_front"))
        self.assertLess(q.bounds[3], 1.0)                          # on the open front, outside the well

    def test_a_lift_open_on_three_sides_has_no_door_found(self):
        lift = box(0, 0, 2500, 2500)
        walls = box(-200, 2500, 2700, 2700)                        # walled on one side only
        self.assertEqual(el.lift_door_quads([lift], walls, open_front=True), [])

    def test_which_spaces_are_lifts(self):
        self.assertTrue(el.is_lift_room("24 PAX P.LIFT", "Lift"))
        self.assertFalse(el.is_lift_room("LIFT PRE", "Lift"))
        self.assertFalse(el.is_lift_room("LIFT WELL BELOW", "Lift"))
        self.assertFalse(el.is_lift_room("LIFT LOBBY", "Lift Lobby"))


class CutoutRunTests(unittest.TestCase):

    def test_in_a_wall_run_or_lining_a_face(self):
        walls = unary_union([box(0, 0, 2000, 200), box(2600, 0, 5000, 200)])
        self.assertTrue(el.cutout_in_wall_run(box(2000, 0, 2600, 200), walls))
        self.assertFalse(el.cutout_in_wall_run(box(1000, 200, 1600, 500), walls))


class FamilyTests(unittest.TestCase):

    def test_lift_pre_is_a_shaft(self):
        fr = dr._FAMILY_RULES
        for n in ("LIFT PRE", "LIFT PRESS", "LIFT PRESSURISATION"):
            self.assertEqual(dr.family_for_label(n, fr)["name"], "Shaft", n)
        for n in ("LIFT WELL BELOW", "24 PAX P.LIFT", "SERVICE LIFT"):
            self.assertEqual(dr.family_for_label(n, fr)["name"], "Lift", n)


class ShaftCellTests(unittest.TestCase):

    def test_the_shaft_runs_to_its_door(self):
        class T:
            def __init__(self, g):
                self.ok, self.polygon, self.area_sqm = True, list(g.exterior.coords)[:-1], g.area / 1e6
        cell = box(0, 300, 600, 900)

        def trace(x, y):
            return T(cell) if cell.contains(__import__("shapely.geometry", fromlist=["Point"]).Point(x, y)) else None
        merged = [{"bbox": (0, 0, 600, 300), "poly": rect(0, 0, 600, 300), "holes": [], "label": "TEN",
                   "source": "X", "cx": 300, "cy": 150, "sr_no": 1, "height": 100.0}]
        doors = [rect(0, 900, 600, 1100)]
        n = dr._shaft_cells_to_door(merged, [], doors, lambda d: Polygon(d["poly"]), lambda d: "MEP / Services",
                                    trace, lambda: None, lambda t, x, y: None)
        self.assertEqual(n, 1)
        new = merged[-1]
        self.assertEqual(new["source"], "S")
        self.assertEqual(new["label"], "TEN")
        self.assertAlmostEqual(Polygon(new["poly"]).area / 1e6, 0.36, places=2)

    def test_no_door_no_room(self):
        class T:
            def __init__(self, g):
                self.ok, self.polygon, self.area_sqm = True, list(g.exterior.coords)[:-1], g.area / 1e6
        cell = box(0, 300, 600, 900)
        from shapely.geometry import Point
        merged = [{"bbox": (0, 0, 600, 300), "poly": rect(0, 0, 600, 300), "holes": [], "label": "TEN",
                   "source": "X", "cx": 300, "cy": 150, "sr_no": 1, "height": 100.0}]
        n = dr._shaft_cells_to_door(merged, [], [rect(5000, 5000, 5600, 5200)], lambda d: Polygon(d["poly"]),
                                    lambda d: "MEP / Services",
                                    lambda x, y: T(cell) if cell.contains(Point(x, y)) else None,
                                    lambda: None, lambda t, x, y: None)
        self.assertEqual(n, 0)


class StatementOwnerTests(unittest.TestCase):
    """XA: an X strip of the ELE SH shaft stands behind ELECTRICAL SH's wall;
    it is ELE SH's opening, so ELECTRICAL SH stays a deducted shaft."""

    def test_an_opening_behind_the_wall_along_another_shaft_is_that_ones(self):
        from Drawing_Analyzer.area_statement import classify
        nb = [{"name": "Unknown Shaft", "opening": True, "inside": False, "shared_mm": 0.0, "touches": ["ELE SH"]}]
        self.assertEqual(classify("ELECTRICAL SH", "Shaft", opening=False, neighbours=nb), "elect_shaft")

    def test_an_opening_along_the_compartment_is_its_own(self):
        from Drawing_Analyzer.area_statement import classify
        nb = [{"name": "Unknown Shaft", "opening": True, "inside": False, "shared_mm": 900.0, "touches": ["ELE SH"]}]
        self.assertEqual(classify("ELECTRICAL SH", "Shaft", opening=False, neighbours=nb), "shaft_room")
        nb = [{"name": "Unknown Shaft", "opening": True, "inside": False, "shared_mm": 0.0, "touches": []}]
        self.assertEqual(classify("ELECTRICAL SH", "Shaft", opening=False, neighbours=nb), "shaft_room")


class BalconyEndWallTests(unittest.TestCase):

    def test_the_balcony_grows_over_its_parapet_not_its_end_wall(self):
        plate = rect(0, 0, 30000, 30000)
        bal = {"bbox": (26000, 200, 29800, 2200), "poly": rect(26000, 200, 29800, 2200), "holes": [],
               "label": "BALCONY", "source": "P", "cx": 27900, "cy": 1200, "sr_no": 1}
        parapet = box(25800, 0, 30000, 200)
        end_wall = box(29800, 200, 30000, 2400)
        shapes = {}

        def shape_of(d):
            return shapes.get(id(d)) or Polygon(d["poly"])

        def set_shape(d, g, c):
            shapes[id(d)] = g
            return True
        n = dr._grow_balconies([bal], {LN.WALL: [parapet, end_wall]}, [], plate, shape_of, set_shape,
                               dr.family_for_label)
        self.assertEqual(n, 1)
        g = shapes[id(bal)]
        self.assertGreater(g.intersection(parapet).area, 0.5 * parapet.area)
        self.assertLess(g.intersection(end_wall).area, 0.05 * end_wall.area)


class DoorLayerEvidenceTests(unittest.TestCase):
    """Run on BHUB first floor and H10 Tower 1 1st floor: what is drawn on
    the door layer, or across an opening, is read for what it is."""

    def facts(self, o, segs, dashed=(), texts=(), walls=None):
        ev = E.Evidence(segs, dashed, (), texts)
        f = E._opening_facts(o, ev, walls)
        return f, {E.DOOR: E.score_door(f)}

    def test_treads_drawn_on_the_door_layer_are_flights(self):
        """H10 Tower 1 draws its stair treads on A3-DOOR."""
        stair = box(0, 0, 4500, 7000)
        segs = treads(0, 2000, 1500, 13, layer="A3-DOOR") + treads(2500, 4500, 1500, 13, layer="A3-DOOR")
        segs = [s[:7] + (True,) + s[8:] for s in segs]
        fl = E.flights_in(stair, E.Evidence(segs))
        self.assertEqual(sorted(f["n"] for f in fl), [13, 13])
        sp = [{"poly": stair, "name": "STAIRCASE", "family": "Staircase", "current_class": E.STAIR}]
        recs = E.identify_floor(sp, [], None, E.Evidence(segs), floor="F", prefix="F1")
        st = [r for r in recs if r["name"] == "STAIRCASE"][0]
        self.assertEqual(st["classification"], E.STAIR)
        self.assertTrue(st["agrees"])

    def test_a_few_door_layer_lines_are_no_flight(self):
        segs = [s[:7] + (True,) + s[8:] for s in treads(0, 900, 0, 5, layer="A3-DOOR")]
        self.assertEqual(E.flights_in(box(0, -200, 1500, 2000), E.Evidence(segs)), [])

    def test_treads_on_the_stair_layer_come_first(self):
        stair = box(0, 0, 4500, 7000)
        segs = treads(0, 2000, 1500, 12) + [s[:7] + (True,) + s[8:]
                                            for s in treads(2500, 4500, 1500, 13, layer="A3-DOOR")]
        self.assertEqual([f["n"] for f in E.flights_in(stair, E.Evidence(segs))], [12])

    def test_a_cut_outs_ends_on_the_door_layer_are_its_jambs_not_a_leaf(self):
        """H10 Tower 1: the wall is drawn again on A3-DOOR, so the X-marked
        cut-out's 300 mm ends are door-layer lines. They are its jambs."""
        walls = wall_with_gap(1000, 1700, t=300.0)
        cut = box(1000, 0, 1700, 300)
        mark = outline(1000, 0, 1700, 300, "A3-CUTOUTS") + x_mark(1000, 0, 1700, 300, "A3-CUTOUTS")
        jambs = [seg(1000, 0, 1000, 300, "A3-DOOR", True), seg(1700, 300, 1700, 0, "A3-DOOR", True)]
        f, sc = self.facts({"poly": cut, "kind": "cutout", "x": True}, mark + jambs, dashed=keys(mark), walls=walls)
        self.assertFalse(f["leaf"])
        self.assertEqual(f["boundary"], "dashed")

    def test_an_open_leaf_at_the_jamb_is_still_a_leaf(self):
        walls = wall_with_gap(1000, 1900)
        segs = [seg(1000, 20, 1000, 180, "A-DOOR", True), seg(1900, 20, 1900, 180, "A-DOOR", True),
                seg(1000, 200, 1000, 1100, "A-DOOR", True)]
        f, _sc = self.facts({"poly": box(1000, 0, 1900, 200), "kind": "door"}, segs, walls=walls)
        self.assertTrue(f["leaf"])

    def test_a_louvres_lines_jamb_to_jamb_are_glazing_not_sliding_panels(self):
        """BHUB first floor: a 1200 mm louvre (LUV-3, W-3) in the outer wall,
        its lines on EL-SHAFT. A door's tag 400 mm off does not make it a door."""
        walls = wall_with_gap(1000, 2200)
        quad = box(1000, 0, 2200, 200)
        segs = [seg(1000, 0, 1000, 200, "A3-WL_EXT"), seg(2200, 0, 2200, 200, "A3-WL_EXT")] + \
            [seg(1000, y, 2200, y, "EL-SHAFT") for y in (0, 75, 125, 200)]
        texts = [{"text": "W-3", "x": 1600, "y": 500}, {"text": "D-4", "x": 2600, "y": -300}]
        f, sc = self.facts({"poly": quad, "kind": "window"}, segs, texts=texts, walls=walls)
        self.assertEqual(f["panels"], 0)
        self.assertFalse(f["sliding"])
        self.assertTrue(f["window_tag"])

    def test_staggered_panels_still_slide(self):
        walls = wall_with_gap(1000, 2200)
        segs = [seg(1000, 0, 1000, 200, "A-WALL"), seg(2200, 0, 2200, 200, "A-WALL"),
                seg(1000, 80, 1700, 80, "A-FURN"), seg(1500, 120, 2200, 120, "A-FURN")]
        f, _sc = self.facts({"poly": box(1000, 0, 2200, 200), "kind": "door"}, segs, walls=walls)
        self.assertEqual(f["panels"], 2)
        self.assertTrue(f["sliding"])


if __name__ == "__main__":
    unittest.main()
