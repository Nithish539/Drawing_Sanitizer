# -*- coding: utf-8 -*-
"""Area Statement: column allocation, floor titles, towers and the workbook.

Run:  python -m unittest discover -s tests -v
"""

import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from openpyxl import load_workbook  # noqa: E402

from Drawing_Analyzer.area_statement import (build_statement,  # noqa: E402
                                             classify, level_name,
                                             write_area_statement)
from Drawing_Analyzer.levels import (floor_name, parse_levels,  # noqa: E402
                                     tower_from_name, tower_of)


class ClassifyTests(unittest.TestCase):
    """Names as they come off the A-106..A-114 and T1 drawings."""

    CASES = [
        ("GENTS TOILET", "Toilet", "toilets"),
        ("LADIES TOILET", "Toilet", "toilets"),
        ("Janit.", "Unknown", "toilets"),
        ("PCT", "Unknown", "toilets"),
        ("Toilet Exhaust", "Toilet", "toilet_exhaust"),
        ("TOI.EX.SHAFT", "Shaft", "toilet_exhaust"),
        ("P.Lift LOBBY {HIGH RISE}", "Lift", "lift_lobby"),
        ("LOBBY PREZZ.", "Circulation", "lobby_press"),
        ("Lobby Press. shaft", "Shaft", "lobby_press"),
        ("LIFT PRE", "Lift", "lobby_press"),
        ("STAIR PREZ", "Staircase", "stair_press"),
        ("STAIR PR", "Staircase", "stair_press"),
        ("STAIRCASE PRE.SHAFT", "Staircase", "stair_press"),
        ("PRESS. SHAFT", "Shaft", "stair_press"),
        ("PRESS", "Unknown", "stair_press"),
        ("AHU ROOM", "MEP / Services", "ahu_room"),
        ("AHU PIPE SHAFT", "Shaft", "ahu_shaft"),
        ("CHW,TFA & RWPs SHAFT", "Shaft", "chw"),
        ("ELECTRICAL ROOM", "Electrical", "ele_room"),
        ("Elec.", "Electrical", "elect_shaft"),
        ("ELE SHAFT", "Shaft", "elect_shaft"),
        ("BALCONY 100MM SUNK", "Balcony / Open", "balcony"),
        ("REFUGE AREA-1", "Unknown", "refuge"),
        ("STAIRCASE", "Staircase", "staircase"),
        ("ST-01", "Staircase", "staircase"),
        # a fire tower is the fireman's lobby: a service lobby, never a staircase
        ("FIRE TOWER-1", "Fire / Life Safety", "service_lobby"),
        ("P.LIFT/FIRE LIFT-12", "Lift", "lift_shaft"),
        ("24 PAX P.LIFT 2500 X 2500", "Lift", "lift_shaft"),
        ("DG EXHAUST", "Shaft", "dg"),
        ("PHE SHAFT", "Shaft", "phe"),
        ("RWP", "Unknown", "rwdp"),
        ("STP EXHAUST", "Plumbing", "stp"),
        ("COMM.", "MEP / Services", "comm"),
        ("FAPA", "Unknown", "fapa"),
        ("TEN", "Unknown", "ten"),
        ("FHR-3", "Fire / Life Safety", "fhr"),
        ("Unknown Shaft", "Shaft", "unnamed_opening"),
        ("Pantry", "Storage", None),
        ("OFFICE AREA", "Habitable Room", None),
    ]

    def test_openings_against_rooms(self):
        """What the drawing says about how a boundary was found decides
        whether it is a slab opening (deducted) or a room."""
        # an X-marked opening: deducted under its name
        self.assertEqual(classify("PHE SHAFT", "Shaft", opening=True), "phe")
        # a walled shaft room with X-marked openings in it: the openings are
        # deducted, the room is listed
        inside = [{"name": "Unknown Shaft", "layer": "AI_UNKNOWN", "opening": True}]
        self.assertEqual(classify("PHE-SHAFT 2300X1270", "Shaft", opening=False, neighbours=inside), "shaft_room")
        # a walled shaft compartment with no opening drawn is the shaft itself
        self.assertEqual(classify("AHU PIPE SHAFT", "Shaft", opening=False, neighbours=[]), "ahu_shaft")
        # a fire hose cabinet is equipment, not a shaft
        self.assertIsNone(classify("FHC-3", "Fire / Life Safety", opening=False))
        # an X tagged with a column mark or a door tag: an unnamed cut-out
        self.assertEqual(classify("C-7", "Unknown", opening=True), "unnamed_opening")
        self.assertEqual(classify("SD-6", "Unknown", opening=True), "unnamed_opening")
        self.assertEqual(classify("Unknown Shaft", "Unknown", opening=True), "unnamed_opening")
        # the well inside a staircase or lift is part of it
        self.assertIsNone(classify("Unknown Shaft", "Unknown", opening=True, within_family="Staircase"))
        # a lift well found by its X is a lift shaft (vertical), not a cut-out
        self.assertEqual(classify("LIFT WELL 2650X2650", "Lift", opening=True), "lift_shaft")

    def test_lobby_by_its_neighbours(self):
        service = [{"name": "SERVICE LIFT", "layer": "AI_LIFT"}]
        passenger = [{"name": "24 PAX P.LIFT", "layer": "AI_LIFT"}, {"name": "SERVICE LIFT", "layer": "AI_LIFT"}]
        self.assertEqual(classify("LOBBY", "Circulation", opening=False, neighbours=service), "service_lobby")
        self.assertEqual(classify("LOBBY", "Circulation", opening=False, neighbours=passenger), "lift_lobby")
        self.assertEqual(classify("PASSAGE 2265 WIDE", "Circulation", opening=False, neighbours=[]), "circulation")
        self.assertEqual(classify("LIFT LOBBY", "Lift Lobby", opening=False), "lift_lobby")

    def test_names_file_into_the_reference_columns(self):
        for name, family, want in self.CASES:
            with self.subTest(name=name):
                self.assertEqual(classify(name, family), want)


class LevelTests(unittest.TestCase):
    """Plan titles as written under the T1 plates."""

    def labels(self, title):
        return [lv["label"] for lv in parse_levels(title)]

    def test_titles(self):
        self.assertEqual(self.labels("%%U1ST FLOOR PLAN"), ["1st Floor"])
        self.assertEqual(self.labels("%%UTYPICAL FLOOR PLAN-(2nd,4th,6th & 10th)"),
                         ["2nd Floor", "4th Floor", "6th Floor", "10th Floor"])
        self.assertEqual(self.labels("TYPICAL FLOOR PLAN 5TH TO 9TH"),
                         ["5th Floor", "6th Floor", "7th Floor", "8th Floor",
                          "9th Floor"])
        self.assertEqual(self.labels("SECOND & THIRD FLOOR PLAN"),
                         ["2nd Floor", "3rd Floor"])
        self.assertEqual(self.labels("%%UTERRACE FLOOR"), ["Terrace Floor"])
        self.assertEqual(self.labels("GROUND FLOOR PLAN"), ["Ground Floor"])
        self.assertEqual(self.labels("BASEMENT-2 PLAN"), ["Basement 2"])
        self.assertEqual(self.labels("TERRACE BELOW"), [])
        self.assertEqual(self.labels("STRUCTURAL GLAZING TYPE-1"), [])
        self.assertTrue(parse_levels("%%UTYPICAL REFUGE FLOOR PLAN- 5th")[0]["refuge"])

    def test_towers(self):
        self.assertEqual(tower_of("%%UTOWER-1"), "TOWER 1")
        self.assertIsNone(tower_of("SCHEDULE OF OPENINGS(TOWERS)"))
        self.assertEqual(tower_from_name("T1_TYPICAL FLOOR PLANS.dxf"), "TOWER 1")
        self.assertIsNone(tower_from_name("A-109 7th FLOOR.dxf"))

    def test_level_from_the_file_name(self):
        self.assertEqual(level_name("A-109 7th FLOOR (OFFICE FLOOR 04).dxf"),
                         "7th FLOOR")


def _room(name, family, area, floor):
    return {"name": name, "family": family, "area_sqm": area, "floor": floor}


class FloorNameTests(unittest.TestCase):
    """Step 2 names each plate after its plan title."""

    CASES = [
        ("1ST FLOOR PLAN", "1st Floor"),
        ("TYPICAL FLOOR PLAN-(2nd,4th,6th & 10th)",
         "Typical 2nd, 4th, 6th & 10th Floor"),
        ("TYPICAL FLOOR PLAN-(13th,14th,15th & 17th )",
         "Typical 13th, 14th, 15th & 17th Floor"),
        ("TYPICAL FLOOR PLAN 5TH TO 9TH", "Typical 5th to 9th Floor"),
        ("TYPICAL REFUGE FLOOR PLAN- 5th", "Typical 5th Floor (Refuge)"),
        ("12TH REFUGE FLOOR PLAN", "12th Floor (Refuge)"),
        ("SECOND & THIRD FLOOR PLAN", "Typical 2nd & 3rd Floor"),
        ("TERRACE FLOOR", "Terrace Floor"),
        ("UPPER TERRACE FLOOR", "Upper Terrace Floor"),
        ("6th FLOOR PLAN (REFUGE / OFFICE - WING-B) SCALE : 1:150",
         "6th Floor (Refuge) - Wing B"),
        ("A-106 4th FLOOR PLAN(OFFICE FLOOR 01)", "4th Floor"),
        ("GENERAL NOTES", ""),
    ]

    def test_titles(self):
        for title, want in self.CASES:
            with self.subTest(title=title):
                self.assertEqual(floor_name(title), want)


class StatementTests(unittest.TestCase):

    def area(self):
        return {
            "floors": [
                {"name": "P1", "area_sqm": 1000.0, "tower": "TOWER 1",
                 "title": "1ST FLOOR PLAN",
                 "levels": parse_levels("1ST FLOOR PLAN")},
                {"name": "P2", "area_sqm": 1000.0, "tower": "TOWER 1",
                 "title": "TYPICAL FLOOR PLAN-(2nd,4th & 3rd)",
                 "levels": parse_levels("TYPICAL FLOOR PLAN-(2nd,4th & 3rd)")},
                {"name": "P3", "area_sqm": 800.0, "tower": "TOWER 2",
                 "title": "1ST FLOOR PLAN",
                 "levels": parse_levels("1ST FLOOR PLAN")},
            ],
            "rooms": [
                _room("GENTS TOILET", "Toilet", 50.0, "P1"),
                _room("STAIRCASE", "Staircase", 20.0, "P1"),
                _room("P.LIFT-1", "Lift", 10.0, "P1"),
                _room("PHE SHAFT", "Shaft", 4.0, "P1"),
                _room("LADIES TOILET", "Toilet", 45.0, "P2"),
                _room("STAIRCASE", "Staircase", 20.0, "P2"),
                _room("LADIES TOILET", "Toilet", 40.0, "P3"),
            ],
        }

    def test_floor_by_floor_per_tower(self):
        data = build_statement(self.area(), "T1 plans.dxf")
        self.assertEqual([t["title"] for t in data["towers"]],
                         ["TOWER 1", "TOWER 2"])
        t1 = data["towers"][0]["rows"]
        # The typical plate is repeated for each floor it names, in order.
        self.assertEqual([r["level"] for r in t1],
                         ["1st Floor", "2nd Floor", "3rd Floor", "4th Floor"])
        self.assertEqual([r["values"]["toilets"] for r in t1],
                         [50.0, 45.0, 45.0, 45.0])

    def test_same_floor_twice_without_a_tower_is_two_towers(self):
        area = self.area()
        for f in area["floors"]:
            f["tower"] = None
        data = build_statement(area, "plans.dxf")
        self.assertEqual([t["title"] for t in data["towers"]],
                         ["TOWER 1", "TOWER 2"])

    def test_workbook(self):
        """The CLEAN FORMAT workbook: one sheet, every figure a formula over
        the register through defined names."""
        out = os.path.join(tempfile.mkdtemp(), "st.xlsx")
        data = write_area_statement(self.area(), out, "T1 plans.dxf",
                                    project_name="Phoenix - 14")
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Area Statement"])
        ws = wb["Area Statement"]
        self.assertEqual(ws["B2"].value, '="Project Name: "&Project')
        names = set(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else \
            {d.name for d in wb.defined_names.definedName}
        for nm in ("Project", "Drawing", "Tower", "SaleFactor", "SqftPerSqm", "F_Floor", "F_Plate", "F_Order",
                   "M_Key", "M_Header", "M_Group", "M_Deducted", "M_PlinthLess", "G_Group", "G_Title",
                   "R_Name", "R_Layer", "R_Nos", "R_Area", "R_Plate", "R_Key", "LkFloor", "LkElement",
                   "LkCategory", "LayPlate"):
            self.assertIn(nm, names)
        cells = [c for row in ws.iter_rows() for c in row if c.value]
        text = [str(c.value) for c in cells]
        self.assertIn("TOWER 1", text)
        self.assertIn("TOWER 2", text)
        # Unused columns are left out; the keys name the columns.
        keys = data["statement_keys"]
        for gone in ("dg", "slab", "elevation", "bms"):
            self.assertNotIn(gone, keys)
        for kept in ("toilets", "staircase", "lift_shaft", "phe", "net_bua", "plinth"):
            self.assertIn(kept, keys)
        # every statement figure is a formula over the register
        krow = next(c.row for c in cells if c.value == "gross" and c.column == 4)
        first = ws.cell(row=krow + 1, column=4)
        self.assertTrue(str(first.value).startswith("=SUMIFS(R_Area,R_Plate,"))
        net_col = 4 + keys.index("net_bua")
        self.assertTrue(str(ws.cell(row=krow + 1, column=net_col).value).startswith("="))
        self.assertNotIn("Room allocation", wb.sheetnames)
        # the register: elements in yellow, keys beside them, plates once
        reg = [c for c in cells if c.column == 9 and c.value in ("toilets", "staircase", "phe")]
        self.assertTrue(reg)
        plates = [c for c in cells if isinstance(c.value, str) and c.value.startswith("PLATE:")]
        self.assertEqual(len(plates), len({c.value for c in plates}))
        # no reconciliation against anything
        self.assertFalse(any("in statement?" in t.lower() or "reconcil" in t.lower() for t in text))


def _col(headers, name):
    from openpyxl.utils import get_column_letter
    return get_column_letter(2 + headers.index(name))



class FamilyFirstTests(unittest.TestCase):
    """A room's family never puts it in a column its name contradicts."""

    def test_fire_tower_never_in_staircase(self):
        self.assertEqual(classify("FIRE TOWER 3100X4400", "Fire / Life Safety"), "service_lobby")
        self.assertNotEqual(classify("FIRE TOWER", "Fire / Life Safety"), "staircase")
        self.assertNotEqual(classify("STAIR FIRE TOWER", "Fire / Life Safety"), "staircase")

    def test_stair_prez_of_family_shaft(self):
        self.assertEqual(classify("STAIR PREZ", "Shaft"), "stair_press")
        self.assertEqual(classify("STAIR PR", "Shaft"), "stair_press")

    def test_staircase_names(self):
        self.assertEqual(classify("ST-02", "Staircase"), "staircase")
        self.assertEqual(classify("STAIRCASE", "Unknown"), "staircase")
        self.assertEqual(classify("STAIRCASE", ""), "staircase")

    def test_name_picks_between_the_familys_columns(self):
        self.assertEqual(classify("DG EXHAUST", "Shaft"), "dg")
        self.assertEqual(classify("LIFT PRE", "Lift"), "lobby_press")
        self.assertEqual(classify("24 PAX P.LIFT 2500 X 2500", "Lift"), "lift_shaft")
        self.assertEqual(classify("LIFT WELL", "Lift"), "lift_shaft")

    def test_statement_and_allocation(self):
        rooms = [_room("STAIRCASE", "Staircase", 32.56, "F1"), _room("STAIRCASE", "Staircase", 32.56, "F1"),
                 _room("FIRE TOWER 3100X4400", "Fire / Life Safety", 13.64, "F1"),
                 _room("FIRE TOWER 3100X4400", "Fire / Life Safety", 13.64, "F1"),
                 _room("SHAFT_06", "Shaft", 1.73, "F1")]
        data = {"floors": [{"name": "F1", "area_sqm": 2000.0, "levels": [{"key": 1, "label": "1st Floor"}]}],
                "rooms": rooms}
        st = build_statement(data, "x.dxf")
        v = st["towers"][0]["rows"][0]["values"]
        self.assertAlmostEqual(v["staircase"], 65.12, places=2)
        self.assertAlmostEqual(v["service_lobby"], 27.28, places=2)
        self.assertAlmostEqual(v["unnamed_opening"], 1.73, places=2)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "s.xlsx")
            write_area_statement(data, out, "x.dxf")
            ws = load_workbook(out)["Area Statement"]
            # the register files the shaft under its deducted key, the
            # staircases under theirs
            reg = {}
            for row in ws.iter_rows(min_row=20, values_only=True):
                if row[2] and row[8] in ("unnamed_opening", "staircase", "service_lobby"):
                    reg.setdefault(row[8], []).append((row[2], row[4], row[5]))
        self.assertEqual(reg["unnamed_opening"], [("SHAFT_06", 1, 1.73)])
        self.assertEqual(reg["staircase"], [("STAIRCASE", 2, 65.12)])
        self.assertEqual(reg["service_lobby"][0][1:], (2, 27.28))

    def test_template_composes_columns(self):
        cfg = {"area_statement": {"templates": {"t": {"columns": [
            {"key": "ahu_room", "header": "AHU shafts", "sum": ["ahu_room", "ahu_shaft"]}]}}}}
        rooms = [_room("AHU ROOM", "MEP / Services", 30.0, "F1"), _room("AHU PIPE SHAFT", "Shaft", 2.0, "F1")]
        data = {"floors": [{"name": "F1", "area_sqm": 1000.0, "levels": [{"key": 1, "label": "1st Floor"}]}],
                "rooms": rooms}
        st = build_statement(data, "x.dxf", cfg=cfg, template="t")
        row = st["towers"][0]["rows"][0]
        self.assertAlmostEqual(row["values"]["ahu_room"], 32.0)
        self.assertAlmostEqual(row["canon"]["ahu_shaft"], 2.0)      # canonical figures stay apart
        self.assertEqual(st["template"]["header_of"]["ahu_room"], "AHU shafts")


if __name__ == "__main__":
    unittest.main()
