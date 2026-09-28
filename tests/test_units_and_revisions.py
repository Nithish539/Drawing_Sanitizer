# -*- coding: utf-8 -*-
"""Tests for unit resolution and revision-file loading.

The unit bug these cover was not a rounding error: reading a millimetre
drawing as metres multiplied every area by a million, which then let
every stray polyline clear the floor-plate filter and pushed every wall
past the ray-cast search radius, so step 3 drew nothing at all. One
wrong setting, three symptoms.

Run:  python -m unittest discover -s tests -v
"""

import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from Drawing_Analyzer.revisions import _parse_sheet, load_revision  # noqa: E402
from Drawing_Analyzer.units import (check_plausibility,  # noqa: E402
                                    infer_units_from_spans, resolve_units,
                                    span_stats, units_from_code)

MM, CM, M, FT, INCH = 4, 5, 6, 2, 1


def spans_for(building_m, count=120, unit_mm=1.0):
    """Synthesise a plausible floor-plan size distribution, in the given
    unit: lots of small elements, a few room-sized ones."""
    scale = 1000.0 / unit_mm
    out = [0.9 * scale] * (count // 2)          # 0.9 m elements
    out += [3.0 * scale] * (count // 3)         # 3 m rooms
    out += [building_m * 0.4 * scale] * 6       # wings
    out += [building_m * scale]                 # the plate
    return out


class TestUnitInference(unittest.TestCase):
    """The unit is read from the size distribution, not the header."""

    def test_millimetre_drawing(self):
        u = infer_units_from_spans(spans_for(72, unit_mm=1.0))
        self.assertEqual(u.mm_per_unit, 1.0)

    def test_metre_drawing(self):
        u = infer_units_from_spans(spans_for(72, unit_mm=1000.0))
        self.assertEqual(u.mm_per_unit, 1000.0)

    def test_centimetre_drawing(self):
        u = infer_units_from_spans(spans_for(72, unit_mm=10.0))
        self.assertEqual(u.mm_per_unit, 10.0)

    def test_nothing_plausible(self):
        self.assertIsNone(infer_units_from_spans([]))
        self.assertIsNone(infer_units_from_spans([1e-9, 2e-9]))

    def test_outliers_do_not_swing_it(self):
        """A title block parked far away must not decide the unit.

        The real file spans 1,224,573 units for that reason while the
        building is 72 m across, which is why a single drawing-wide
        extent cannot be used.
        """
        spans = spans_for(72, unit_mm=1.0) + [1224573.0, 900000.0]
        self.assertEqual(infer_units_from_spans(spans).mm_per_unit, 1.0)


class TestHeaderOverriding(unittest.TestCase):
    """A DXF may simply declare the wrong unit."""

    def setUp(self):
        # $INSUNITS says Metres; the geometry says millimetres. This is
        # the actual file that produced 5,207,904,569 sq.m.
        self.spans = spans_for(72, unit_mm=1.0)

    def test_geometry_beats_a_wrong_header(self):
        units, info = resolve_units(6, self.spans, None)
        self.assertEqual(units.mm_per_unit, 1.0)
        self.assertEqual(units.source, "geometry")
        self.assertIn("header declares Metres", units.note)

    def test_correct_header_is_kept(self):
        units, info = resolve_units(4, self.spans, None)
        self.assertEqual(units.mm_per_unit, 1.0)
        self.assertEqual(units.source, "header")

    def test_override_wins_over_both(self):
        units, info = resolve_units(6, self.spans, "6")
        self.assertEqual(units.mm_per_unit, 1000.0)
        self.assertEqual(units.source, "override")
        self.assertFalse(info["plausible"])
        self.assertIn("not a building", info["message"])

    def test_no_header_falls_back_to_geometry(self):
        units, _ = resolve_units(0, self.spans, None)
        self.assertEqual(units.mm_per_unit, 1.0)
        self.assertEqual(units.source, "geometry")

    def test_no_geometry_at_all_uses_the_header(self):
        units, _ = resolve_units(6, [], None)
        self.assertEqual(units.mm_per_unit, 1000.0)

    def test_plausible_header_is_not_second_guessed(self):
        """T1_TYPICAL FLOOR PLANS: declared millimetres and drawn in them,
        but full of small hatch squares - median element 0.20 m, just under
        the band. Centimetres scored better and used to win, reading every
        wall ten times too thick."""
        spans = [200.0] * 600 + [900.0] * 100 + [5325.0] * 60 + [72000.0] * 4
        units, _ = resolve_units(4, spans, None)
        self.assertEqual(units.mm_per_unit, 1.0)
        self.assertEqual(units.source, "header")

    def test_area_error_is_a_million(self):
        """Why this matters: the mistake is 10^6 on area."""
        from Drawing_Analyzer.geometry import shoelace_area_sqm
        pts = [(0, 0), (60000, 0), (60000, 40000), (0, 40000)]
        mm = units_from_code(4)
        m = units_from_code(6)
        a_mm = shoelace_area_sqm([mm.point_to_mm(x, y) for x, y in pts])
        a_m = shoelace_area_sqm([m.point_to_mm(x, y) for x, y in pts])
        self.assertAlmostEqual(a_mm, 2400.0, places=3)
        self.assertAlmostEqual(a_m / a_mm, 1_000_000.0, places=0)


class TestPlausibility(unittest.TestCase):
    def test_millimetre_drawing_forced_to_metres_is_flagged(self):
        stats = span_stats(spans_for(72, unit_mm=1.0))
        ok, msg = check_plausibility(units_from_code(6), stats)
        self.assertFalse(ok)
        self.assertIn("not a building", msg)

    def test_correct_unit_passes(self):
        stats = span_stats(spans_for(72, unit_mm=1.0))
        ok, msg = check_plausibility(units_from_code(4), stats)
        self.assertTrue(ok)
        self.assertIsNone(msg)

    def test_no_geometry_is_not_flagged(self):
        self.assertTrue(check_plausibility(units_from_code(4), None)[0])


class TestRevisionJson(unittest.TestCase):
    def test_areas_json_round_trip(self):
        data = {"floor_label": "A-FLOOR", "floor_area_sqm": 1000.0,
                "rooms": [{"sr_no": 1, "name": "OFFICE", "area_sqm": 400.0,
                           "family": "Office"}],
                "families": []}
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False,
                                         encoding="utf-8") as fh:
            json.dump(data, fh)
            path = fh.name
        try:
            out, info = load_revision(path, "areas.json")
            self.assertEqual(info["source"], "json")
            self.assertEqual(len(out["rooms"]), 1)
        finally:
            os.unlink(path)

    def test_unrelated_json_is_rejected(self):
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False,
                                         encoding="utf-8") as fh:
            json.dump({"hello": "world"}, fh)
            path = fh.name
        try:
            with self.assertRaises(ValueError):
                load_revision(path, "notes.json")
        finally:
            os.unlink(path)

    def test_unsupported_extension(self):
        with self.assertRaises(ValueError):
            load_revision("/tmp/whatever.txt", "whatever.txt")


class TestRevisionExcel(unittest.TestCase):
    """The workbook this tool writes, and hand-made sheets like it."""

    def _sheet(self, rows):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        return ws

    def test_reads_the_areas_workbook_layout(self):
        # Header is not on row 1: the areas sheet puts a floor-area block
        # above the table, which is why the parser scans for the header.
        ws = self._sheet([
            ["Floor Area"],
            ["A-FLOOR", 1000.0, "Perimeter", 130.0],
            ["Drawing units", "Millimetres"],
            [],
            ["Room Areas"],
            ["Sr.No", "Room Name", "Area (sq.m)", "Perimeter (m)", "Family"],
            [1, "OPEN OFFICE", 399.84, 82.0, "Office"],
            [2, "TOILET", 60.0, 32.0, "Toilet"],
            [None, "Total", 459.84, 114.0, None],
        ])
        parsed = _parse_sheet(ws)
        self.assertEqual(len(parsed["rooms"]), 2)
        self.assertEqual(parsed["rooms"][0]["name"], "OPEN OFFICE")
        self.assertAlmostEqual(parsed["rooms"][0]["area_sqm"], 399.84)
        self.assertEqual(parsed["rooms"][1]["family"], "Toilet")

    def test_total_and_subtotal_rows_are_skipped(self):
        ws = self._sheet([
            ["Room Name", "Area"],
            ["OFFICE", 100.0],
            ["Subtotal - Office", 100.0],
            ["Grand Total", 100.0],
        ])
        parsed = _parse_sheet(ws)
        self.assertEqual([r["name"] for r in parsed["rooms"]], ["OFFICE"])

    def test_columns_matched_by_name_not_position(self):
        ws = self._sheet([
            ["Category", "Area (sq.m)", "Description"],
            ["Office", 25.5, "MEETING 1"],
        ])
        parsed = _parse_sheet(ws)
        self.assertEqual(parsed["rooms"][0]["name"], "MEETING 1")
        self.assertAlmostEqual(parsed["rooms"][0]["area_sqm"], 25.5)
        self.assertEqual(parsed["rooms"][0]["family"], "Office")

    def test_area_with_a_unit_suffix(self):
        ws = self._sheet([
            ["Room", "Area"],
            ["STORE", "42.75 sq.m"],
            ["PANTRY", "1,250.00"],
        ])
        parsed = _parse_sheet(ws)
        self.assertAlmostEqual(parsed["rooms"][0]["area_sqm"], 42.75)
        self.assertAlmostEqual(parsed["rooms"][1]["area_sqm"], 1250.0)

    def test_sheet_without_a_room_table(self):
        ws = self._sheet([["Notes"], ["nothing useful here"]])
        self.assertIsNone(_parse_sheet(ws))

    def test_full_workbook_load_and_compare(self):
        from openpyxl import Workbook

        from Drawing_Analyzer.quantities import compare_revisions

        def make(path, area):
            wb = Workbook()
            ws = wb.active
            ws.title = "Areas"
            ws.append(["Sr.No", "Room Name", "Area (sq.m)", "Family"])
            ws.append([1, "OPEN OFFICE", area, "Office"])
            ws.append([2, "TOILET", 60.0, "Toilet"])
            wb.save(path)

        d = tempfile.mkdtemp()
        p1 = os.path.join(d, "rev_a.xlsx")
        p2 = os.path.join(d, "rev_b.xlsx")
        make(p1, 400.0)
        make(p2, 450.0)
        try:
            old, i1 = load_revision(p1, "rev_a.xlsx")
            new, i2 = load_revision(p2, "rev_b.xlsx")
            self.assertEqual(i1["source"], "xlsx")
            self.assertEqual(i2["rooms"], 2)
            c = compare_revisions(old, new)
            self.assertEqual(c["summary"]["changed_count"], 1)
            self.assertAlmostEqual(c["changed"][0]["delta_sqm"], 50.0)
            self.assertEqual(c["summary"]["added_count"], 0)
        finally:
            os.unlink(p1)
            os.unlink(p2)
            os.rmdir(d)

    def test_excel_and_json_revisions_compare_against_each_other(self):
        """Mixed sources must line up - the common real case."""
        from openpyxl import Workbook

        from Drawing_Analyzer.quantities import compare_revisions

        d = tempfile.mkdtemp()
        xp = os.path.join(d, "old.xlsx")
        jp = os.path.join(d, "new.json")
        wb = Workbook()
        ws = wb.active
        ws.append(["Room Name", "Area (sq.m)", "Family"])
        ws.append(["OPEN OFFICE", 400.0, "Office"])
        wb.save(xp)
        with open(jp, "w", encoding="utf-8") as fh:
            json.dump({"floor_area_sqm": 1000.0, "rooms": [
                {"sr_no": 1, "name": "OPEN OFFICE", "area_sqm": 420.0,
                 "family": "Office"}]}, fh)
        try:
            old, _ = load_revision(xp, "old.xlsx")
            new, _ = load_revision(jp, "new.json")
            c = compare_revisions(old, new)
            self.assertEqual(c["summary"]["changed_count"], 1)
            self.assertAlmostEqual(c["changed"][0]["delta_sqm"], 20.0)
        finally:
            os.unlink(xp)
            os.unlink(jp)
            os.rmdir(d)


if __name__ == "__main__":
    unittest.main(verbosity=2)
