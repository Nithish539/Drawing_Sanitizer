# -*- coding: utf-8 -*-
"""Revision inputs.

The comparison step originally took only the areas JSON this tool
exports. In practice a revision arrives as whatever the last person
saved: the JSON, the areas workbook, or just the corrected DXF. This
normalises all three into the same shape so the comparison itself does
not care where the numbers came from.

  * **JSON**  - an areas export, used as-is.
  * **XLSX**  - the areas workbook, or any sheet with room-name and area
                columns. Headers are matched by name, not position, so a
                reordered or hand-edited sheet still loads.
  * **DXF**   - a corrected drawing; the area calculation is run on it,
                picking the floor-outline layer automatically unless one
                is given.
"""

from __future__ import annotations

import json
import os
from typing import Dict, Optional, Tuple

# Header synonyms, lower-cased and stripped of punctuation. The areas
# workbook this tool writes is the first entry in each list.
_NAME_KEYS = ("room name", "name", "room", "space", "space name",
              "description", "room description")
_AREA_KEYS = ("area sqm", "area sq m", "area", "area m2", "area sqmt",
              "carpet area", "area sq mt", "area sqmts")
_FAMILY_KEYS = ("family", "category", "type", "room type", "group")
_PERIM_KEYS = ("perimeter m", "perimeter", "perimeter mt")
_SR_KEYS = ("sr no", "srno", "s no", "sl no", "sr", "#")


def _norm(text) -> str:
    if text is None:
        return ""
    out = []
    for ch in str(text).strip().lower():
        if ch.isalnum() or ch == " ":
            out.append(ch)
        elif ch in "._-/\\()":
            out.append(" ")
    return " ".join("".join(out).split())


def _match(header, keys) -> bool:
    h = _norm(header)
    if not h:
        return False
    for k in keys:
        if h == k or h.startswith(k + " ") or h.startswith(k):
            return True
    return False


def load_revision(path: str, original_name: str = "",
                  floor_layer: Optional[str] = None,
                  unit_override=None) -> Tuple[Dict, Dict]:
    """Return (areas_dict, info) for a JSON, XLSX or DXF revision file."""
    ext = os.path.splitext(original_name or path)[1].lower()
    if ext == ".json":
        return _from_json(path)
    if ext in (".xlsx", ".xlsm"):
        return _from_xlsx(path)
    if ext == ".dxf":
        return _from_dxf(path, floor_layer, unit_override)
    raise ValueError(
        "Unsupported file type '{0}'. Upload the areas JSON, the areas "
        "Excel workbook, or a corrected DXF.".format(ext or "unknown"))


# -- JSON --------------------------------------------------------------------
def _from_json(path) -> Tuple[Dict, Dict]:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict) or "rooms" not in data:
        raise ValueError("That JSON is not an areas export - it has no "
                         "'rooms' list. Use the 'Download JSON' file from "
                         "the area calculation step.")
    return data, {"source": "json", "rooms": len(data.get("rooms") or [])}


# -- Excel -------------------------------------------------------------------
def _from_xlsx(path) -> Tuple[Dict, Dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        parsed = _parse_sheet(ws)
        if parsed and (best is None or len(parsed["rooms"]) > len(best["rooms"])):
            best = parsed
            best["sheet"] = ws.title
    wb.close()

    if not best or not best["rooms"]:
        raise ValueError(
            "No room table found in that workbook. It needs a header row "
            "with a room-name column and an area column - the 'Areas' sheet "
            "this tool exports already has them.")

    data = {
        "floor_label": best.get("floor_label") or "",
        "floor_area_sqm": best.get("floor_area_sqm"),
        "floor_perimeter_m": None,
        "units": {"label": "from spreadsheet"},
        "rooms": best["rooms"],
        "families": _regroup(best["rooms"]),
    }
    return data, {"source": "xlsx", "sheet": best.get("sheet"),
                  "rooms": len(best["rooms"])}


def _parse_sheet(ws):
    """Find the header row anywhere in the sheet, then read rows under it.

    Scanning for the header rather than assuming row 1 is what lets this
    read the areas workbook, which puts a floor-area block above the
    table, as well as a sheet someone has added a title to.
    """
    rows = []
    for r in ws.iter_rows(max_row=400, values_only=True):
        rows.append(r)
    if not rows:
        return None

    header_at = None
    cols = {}
    for i, row in enumerate(rows):
        found = {}
        for j, cell in enumerate(row):
            if _match(cell, _NAME_KEYS) and "name" not in found:
                found["name"] = j
            elif _match(cell, _AREA_KEYS) and "area" not in found:
                found["area"] = j
            elif _match(cell, _FAMILY_KEYS) and "family" not in found:
                found["family"] = j
            elif _match(cell, _PERIM_KEYS) and "perimeter" not in found:
                found["perimeter"] = j
            elif _match(cell, _SR_KEYS) and "sr" not in found:
                found["sr"] = j
        if "name" in found and "area" in found:
            header_at = i
            cols = found
            break
    if header_at is None:
        return None

    # A "Floor Area" label above the table, if present.
    floor_label, floor_area = "", None
    for row in rows[:header_at]:
        for j, cell in enumerate(row):
            if _norm(cell) == "floor area":
                for k in range(j, min(j + 4, len(row))):
                    v = _to_float(row[k])
                    if v:
                        floor_area = v
                        break
                nxt = rows[rows.index(row) + 1] if rows.index(row) + 1 < len(rows) else ()
                if nxt and nxt[0]:
                    floor_label = str(nxt[0])
                    if floor_area is None:
                        for cellv in nxt[1:4]:
                            v = _to_float(cellv)
                            if v:
                                floor_area = v
                                break

    out = []
    sr = 0
    for row in rows[header_at + 1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        name = row[cols["name"]] if cols["name"] < len(row) else None
        area = _to_float(row[cols["area"]] if cols["area"] < len(row) else None)
        if name is None or area is None:
            continue
        label = str(name).strip()
        if not label or _norm(label) in ("total", "grand total", "subtotal"):
            continue
        if _norm(label).startswith("subtotal"):
            continue
        sr += 1
        rec = {"sr_no": sr, "name": label, "area_sqm": round(area, 2)}
        if "family" in cols and cols["family"] < len(row):
            rec["family"] = str(row[cols["family"]] or "Unknown").strip() \
                or "Unknown"
        else:
            rec["family"] = "Unknown"
        if "perimeter" in cols and cols["perimeter"] < len(row):
            pv = _to_float(row[cols["perimeter"]])
            if pv is not None:
                rec["perimeter_m"] = round(pv, 2)
        out.append(rec)

    return {"rooms": out, "floor_label": floor_label,
            "floor_area_sqm": floor_area}


def _to_float(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    txt = str(v).strip().replace(",", "")
    for suffix in ("sq.m", "sqm", "m2", "m\u00b2", "sq m", "m"):
        if txt.lower().endswith(suffix):
            txt = txt[: -len(suffix)].strip()
            break
    if not txt or txt.startswith("="):
        return None
    try:
        return float(txt)
    except ValueError:
        return None


def _regroup(rooms):
    by = {}
    for r in rooms:
        by.setdefault(r.get("family", "Unknown"), []).append({
            "name": r["name"], "area_sqm": r["area_sqm"],
            "perimeter_m": r.get("perimeter_m", 0.0)})
    out = []
    for name in sorted(by):
        items = sorted(by[name], key=lambda x: -x["area_sqm"])
        out.append({"name": name, "count": len(items),
                    "total_sqm": round(sum(i["area_sqm"] for i in items), 2),
                    "rooms": items})
    return out


# -- DXF ---------------------------------------------------------------------
def _from_dxf(path, floor_layer=None, unit_override=None) -> Tuple[Dict, Dict]:
    from .dxf_reader import compute_room_areas
    from .dxfio import load_document_ex
    from .geometry import shoelace_area_sqm

    doc, msp, units, _info = load_document_ex(path, unit_override)
    loaded = (doc, msp, units)

    if not floor_layer:
        floor_layer = _guess_floor_layer(msp, units, shoelace_area_sqm)
    if not floor_layer:
        raise ValueError(
            "Could not find a floor-outline polyline in that DXF. Run it "
            "through the area calculation step first, or name the floor "
            "layer explicitly.")

    data = compute_room_areas(path, floor_layer, unit_override=unit_override,
                              loaded=loaded)
    return data, {"source": "dxf", "floor_layer": floor_layer,
                  "units": units.label, "rooms": len(data.get("rooms") or [])}


def _guess_floor_layer(msp, units, area_fn):
    """The layer holding the single largest closed polyline.

    That is the floor outline in every drawing this tool produces, and
    picking it automatically means a DXF can be dropped straight into
    the comparison without a second dialogue.
    """
    best_area, best_layer = 0.0, None
    for e in msp.query("LWPOLYLINE POLYLINE"):
        try:
            closed = bool(e.closed) if e.dxftype() == "LWPOLYLINE" \
                else bool(e.is_closed)
            if not closed:
                continue
            if e.dxftype() == "LWPOLYLINE":
                pts = [(p[0], p[1]) for p in e.get_points()]
            else:
                pts = [(v.dxf.location.x, v.dxf.location.y) for v in e.vertices]
        except (AttributeError, TypeError, ValueError):
            continue
        if len(pts) < 3:
            continue
        layer = (e.dxf.layer or "0")
        from . import layernames as LN
        if LN.is_app_layer(layer):
            continue                    # rooms, holes, walls - not the outline
        mm = units.mm_per_unit
        a = area_fn([(x * mm, y * mm) for x, y in pts])
        if a > best_area:
            best_area, best_layer = a, layer
    return best_layer
