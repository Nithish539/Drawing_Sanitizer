# -*- coding: utf-8 -*-
"""Excel writers for the derived reports.

Follows the same conventions as the existing areas workbook: navy headers,
banded borders, and live SUM formulas rather than baked-in totals, so the
client can edit a number and the sheet stays consistent.
"""

from __future__ import annotations

from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(bold=True, size=13, color=NAVY)
SUB_FONT = Font(bold=True, italic=True, size=11)
SUB_FILL = PatternFill("solid", fgColor="F2F5FA")
WARN_FILL = PatternFill("solid", fgColor="FDE7E7")
OK_FILL = PatternFill("solid", fgColor="EAF6EC")
CENTER = Alignment(horizontal="center", vertical="center")
_THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)


def _title(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = TITLE_FONT
    return row + 1


def _kv(ws, row, label, value, bold_value=False):
    ws.cell(row=row, column=1, value=label).font = Font(bold=True)
    c = ws.cell(row=row, column=2, value=value)
    if bold_value:
        c.font = Font(bold=True)
    return row + 1


def _headers(ws, row, headers, widths=None):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return row + 1


def _rows(ws, row, data: Sequence[Sequence], fill_when=None):
    for r in data:
        fill = fill_when(r) if fill_when else None
        for col, v in enumerate(r, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
            if fill:
                c.fill = fill
        row += 1
    return row


def add_validation_sheet(path: str, validation: Dict):
    """Append a Validation sheet to an existing workbook: the checks run on
    the boundaries a Generate wrote (validate.py), then every issue."""
    from openpyxl import load_workbook
    if not validation or validation.get("error"):
        return
    wb = load_workbook(path)
    ws = wb.create_sheet("Validation", 0)
    row = _title(ws, 1, "Output validation")
    row = _kv(ws, row, "Boundaries written", validation.get("boundaries", 0))
    row = _kv(ws, row, "Voids", validation.get("voids", 0))
    row = _kv(ws, row, "Unknown-family boundaries", validation.get("unknown_boundaries", 0))
    row = _kv(ws, row, "Shafts with no name", validation.get("unnamed_shafts", 0))
    row += 1
    row = _headers(ws, row, ["Check", "Issues", "Result"], [42, 12, 12])
    row = _rows(ws, row, [(c["check"], c["count"], "OK" if c["ok"] else "CHECK")
                          for c in validation.get("checks", [])],
                fill_when=lambda r: OK_FILL if r[2] == "OK" else WARN_FILL)
    row += 1
    row = _headers(ws, row, ["Kind", "Severity", "Floor", "Layer", "Label",
                             "Area (sq.m)", "Detail"])
    for i, w in enumerate([12, 10, 22, 22, 40, 12, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _rows(ws, row, [(i["kind"], i["severity"], i["floor"], i["layer"], i["label"],
                     i["area_sqm"], i["detail"]) for i in validation.get("issues", [])],
          fill_when=lambda r: WARN_FILL if r[1] == "error" else None)
    wb.save(path)


def _total(ws, row, label, col_letter, first, last, label_col=2, value_col=3):
    if last < first:
        return row
    c = ws.cell(row=row, column=label_col, value=label)
    c.font = SUB_FONT
    c.fill = SUB_FILL
    v = ws.cell(row=row, column=value_col,
                value="=SUM({0}{1}:{0}{2})".format(col_letter, first, last))
    v.font = SUB_FONT
    v.fill = SUB_FILL
    v.number_format = "0.00"
    return row + 1


# -- layer inventory --------------------------------------------------------
def write_layers_excel(inventory: Dict, out_path: str, drawing_name: str = ""):
    """Two sheets: every layer in one flat table, and the groups summarised.

    Flat first, because the usual reason to open this is to search for a
    layer name and see what the program made of it.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Layers"
    summary = inventory.get("summary") or {}
    row = _title(ws, 1, "Drawing layers" + (" - " + drawing_name
                                            if drawing_name else ""))
    row = _kv(ws, row, "Layers in the file", summary.get("layers", 0))
    row = _kv(ws, row, "Layers carrying content",
              summary.get("layers_with_content", 0))
    row = _kv(ws, row, "Entities in modelspace", summary.get("entities", 0))
    row = _kv(ws, row, "Entities inside blocks",
              summary.get("entities_in_blocks", 0))
    row = _kv(ws, row, "Layers read as walls", summary.get("wall_layers", 0))
    row += 1

    row = _headers(ws, row, [
        "Layer", "Group", "Group is", "Family", "Discipline",
        "Modelspace", "In blocks", "Texts", "Closed polylines",
        "Counts as wall", "Colour", "Lineweight", "Linetype", "State",
        "Mostly"],
        [34, 20, 12, 18, 16, 12, 11, 8, 15, 15, 8, 11, 16, 12, 30])
    first = row
    data = []
    for L in inventory.get("layers", []):
        state = ", ".join([w for w, on in (("off", L.get("off")),
                                           ("frozen", L.get("frozen")),
                                           ("locked", L.get("locked"))) if on])
        data.append([
            L["name"], L["group"], L["group_kind"], L.get("family") or "",
            L.get("discipline") or "", L["entities"], L["in_blocks"],
            L["texts"], L["closed_polylines"], L["counts_as_wall"],
            L["color"], L["lineweight"], L.get("linetype", ""),
            state or "on", L.get("kinds", ""),
        ])
    row = _rows(ws, row, data,
                fill_when=lambda r: (WARN_FILL if r[5] + r[6] == 0 else None))
    ws.freeze_panes = ws.cell(row=first, column=1)

    ws2 = wb.create_sheet("Groups")
    row = _title(ws2, 1, "Layers by group")
    row += 1
    row = _headers(ws2, row, ["Group", "Group is", "Layers", "Entities",
                              "Texts", "Layer names"],
                   [24, 14, 9, 12, 9, 90])
    _rows(ws2, row, [[g["name"], g["kind"], g["layer_count"], g["entities"],
                      g["texts"],
                      ", ".join(L["name"] for L in g["layers"])]
                     for g in inventory.get("groups", [])])
    wb.save(out_path)


# -- quantity report --------------------------------------------------------
def _sheet_title(name, used):
    """A worksheet name Excel accepts: at most 31 characters, none of the
    characters Excel forbids, and unique in the workbook."""
    base = "".join("_" if ch in "[]:*?/\\" else ch for ch in str(name))[:31]
    base = base or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        suffix = " ({0})".format(n)
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _write_sections(ws, row, rep: Dict, plate_label: str = "Floor plate area"):
    """Reconciliation, perimeter and family split for one report - the whole
    sheet or a single floor. Returns the next free row."""
    rec = rep.get("reconciliation")
    if rec:
        row = _title(ws, row, "Area reconciliation")
        row = _headers(ws, row, ["Measure", "Value", "Unit"],
                       [34, 18, 12, 20, 18, 20])
        row = _rows(ws, row, [
            [plate_label, rec["floor_area_sqm"], "sq.m"],
            ["Measured room area", rec["measured_room_area_sqm"], "sq.m"],
            ["Unaccounted (circulation + walls)", rec["unaccounted_sqm"], "sq.m"],
            ["Coverage", rec["coverage_pct"], "%"],
            ["Rooms measured", rec["room_count"], "nos"],
        ])
        ws.cell(row=row, column=1, value=rec["verdict"]).font = Font(italic=True)
        row += 2

    per = rep.get("perimeter")
    if per:
        row = _title(ws, row, "Perimeter")
        row = _headers(ws, row, ["Measure", "Value", "Unit"])
        row = _rows(ws, row, [
            ["Floor outline perimeter", per.get("floor_perimeter_m") or "-", "m"],
            ["Total room perimeter", per["total_room_perimeter_m"], "m"],
            ["Longest room perimeter", per["longest_room_perimeter_m"], "m"],
            ["Shortest room perimeter", per["shortest_room_perimeter_m"], "m"],
        ])
        row += 1

    k = rep.get("kpis")
    if k:
        row = _title(ws, row, "Area split by family")
        row = _headers(ws, row, ["Family", "Count", "Area (sq.m)", "% of floor"])
        first = row
        row = _rows(ws, row, [[f["family"], f["count"], f["total_sqm"],
                               f["pct_of_floor"]] for f in k["family_split"]])
        row = _total(ws, row, "Total", "C", first, row - 1)
        row += 1
    return row


def _write_geometry(ws, row, gh: Dict):
    """The boundary geometry check for one report. Returns the next row."""
    row = _title(ws, row, "Boundary geometry check")
    row = _kv(ws, row, "Boundaries checked", gh["checked"])
    row = _kv(ws, row, "Issues found", gh["issue_count"], bold_value=True)
    row += 1
    if gh["findings"]:
        row = _headers(ws, row, ["Issue", "Boundary A", "Area A", "Boundary B",
                                 "Area B", "Detail"],
                       [18, 30, 12, 30, 12, 70])
        row = _rows(ws, row, [[f["issue"], f["room_a"], f["area_a_sqm"],
                               f["room_b"], f["area_b_sqm"], f["detail"]]
                              for f in gh["findings"]],
                    fill_when=lambda _r: WARN_FILL)
    else:
        ws.cell(row=row, column=1,
                value="No duplicate, contained, sliver or partially "
                      "overlapping boundaries found.").font = Font(italic=True)
        row += 1
    return row + 1


def write_quantities_excel(report: Dict, out_path: str, qa: Dict = None):
    wb = Workbook()
    used = {"summary", "geometry check", "extraction qa"}
    floors = report.get("floors") or []

    # Sheet 1: summary - every floor together
    ws = wb.active
    ws.title = "Summary"
    row = _title(ws, 1, "Derived Quantities and Checks"
                 + (" - all floors" if floors else ""))
    row += 1
    u = report.get("units") or {}
    row = _kv(ws, row, "Floor outline layer", report.get("floor_label") or "-")
    row = _kv(ws, row, "Drawing units", u.get("label", "-"))
    if u.get("assumed"):
        ws.cell(row=row, column=1,
                value="Units were not declared in the DXF header; millimetres "
                      "assumed. Override in the UI if that is wrong.").font = \
            Font(italic=True, color="B06A00")
        row += 1
    row += 1

    if floors:
        # One line per floor up front, so the plates can be compared at a
        # glance; each floor's own sheet follows.
        row = _title(ws, row, "Floor-wise summary")
        row = _headers(ws, row, ["Floor", "Plate area (sq.m)", "Rooms",
                                 "Measured (sq.m)", "Coverage %",
                                 "Geometry issues"],
                       [34, 18, 12, 20, 18, 20])
        first = row
        data = []
        for f in floors:
            rec = f.get("reconciliation") or {}
            gh = f.get("geometry_health") or {}
            data.append([f.get("name"), f.get("area_sqm"),
                         rec.get("room_count", 0),
                         rec.get("measured_room_area_sqm", 0.0),
                         rec.get("coverage_pct", 0.0),
                         gh.get("issue_count", 0)])
        row = _rows(ws, row, data)
        last = row - 1
        for col in ("B", "C", "D"):
            c = ws["{0}{1}".format(col, row)]
            c.value = "=SUM({0}{1}:{0}{2})".format(col, first, last)
            c.font = SUB_FONT
            c.fill = SUB_FILL
        ws.cell(row=row, column=1, value="All floors").font = SUB_FONT
        ws.cell(row=row, column=1).fill = SUB_FILL
        row += 2

    row = _write_sections(ws, row, report,
                          "Floor plate area (all floors)" if floors
                          else "Floor plate area")

    # Sheet 2: geometry health, all floors
    gh = report.get("geometry_health")
    if gh:
        ws2 = wb.create_sheet("Geometry Check")
        _write_geometry(ws2, 1, gh)

    # One sheet per floor: the same checks on that floor's rooms alone.
    for f in floors:
        wsf = wb.create_sheet(_sheet_title(f.get("name") or "Floor", used))
        r = _title(wsf, 1, "Derived Quantities and Checks - "
                   + str(f.get("name") or ""))
        r += 1
        r = _kv(wsf, r, "Floor", f.get("name") or "-")
        r = _kv(wsf, r, "Plate area (sq.m)", f.get("area_sqm"))
        r = _kv(wsf, r, "Drawing units", u.get("label", "-"))
        r += 1
        r = _write_sections(wsf, r, f)
        if f.get("geometry_health"):
            _write_geometry(wsf, r, f["geometry_health"])

    # Sheet 3: extraction QA
    if qa:
        ws3 = wb.create_sheet("Extraction QA")
        r = _title(ws3, 1, "Label to boundary outcome")
        r = _kv(ws3, r, "Boundaries drawn", qa["boundary_count"])
        r = _kv(ws3, r, "Labels with a problem", qa["problem_count"],
                bold_value=True)
        r += 1
        r = _headers(ws3, r, ["Floor", "Sr.No", "Label", "Layer", "Family",
                              "Outcome", "What it means"],
                     [22, 8, 34, 22, 18, 14, 66])
        _rows(ws3, r, [[x["floor"], x["sr_no"], x["text"], x["layer"],
                        x["family"], x["status"], x["meaning"]]
                       for x in qa["rows"]],
              fill_when=lambda row_: None)

    if report.get("cutouts"):
        wc = wb.create_sheet(_sheet_title("Cutouts", used))
        r0 = _title(wc, 1, "Cut-outs in the walls (AI_CUTOUT) - not in any room's area")
        r0 = _headers(wc, r0, ["Floor", "Name", "Area (sq.m)"], [24, 34, 14])
        for c in report["cutouts"]:
            for i, v in enumerate([c.get("floor"), c.get("name"), c.get("area_sqm")], start=1):
                wc.cell(row=r0, column=i, value=v)
            r0 += 1
    wb.save(out_path)


# -- revision comparison ----------------------------------------------------
def write_comparison_excel(cmp_data: Dict, out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = cmp_data["summary"]
    row = _title(ws, 1, "Revision quantity impact")
    row += 1
    row = _headers(ws, row, ["Measure", "Previous", "Current", "Delta"],
                   [34, 18, 18, 18])
    row = _rows(ws, row, [
        ["Floor plate area (sq.m)", s.get("old_floor_area_sqm"),
         s.get("new_floor_area_sqm"), ""],
        ["Total measured area (sq.m)", s["old_total_sqm"], s["new_total_sqm"],
         s["delta_sqm"]],
        ["Change (%)", "", "", s["delta_pct"]],
        ["Room count", s["old_room_count"], s["new_room_count"],
         s["new_room_count"] - s["old_room_count"]],
        ["Rooms added", "", "", s["added_count"]],
        ["Rooms removed", "", "", s["removed_count"]],
        ["Rooms changed", "", "", s["changed_count"]],
    ])

    ws2 = wb.create_sheet("Changed")
    r = _headers(ws2, 1, ["Room", "Family", "Previous (sq.m)", "Current (sq.m)",
                          "Delta (sq.m)", "Delta (%)"],
                 [34, 20, 18, 18, 16, 14])
    _rows(ws2, r, [[x["name"], x["family"], x["old_sqm"], x["new_sqm"],
                    x["delta_sqm"], x["delta_pct"]] for x in cmp_data["changed"]],
          fill_when=lambda row_: WARN_FILL if row_[4] and row_[4] < 0 else OK_FILL)

    ws3 = wb.create_sheet("Added")
    r = _headers(ws3, 1, ["Room", "Family", "Area (sq.m)"], [34, 20, 18])
    _rows(ws3, r, [[x["name"], x["family"], x["new_sqm"]]
                   for x in cmp_data["added"]], fill_when=lambda _r: OK_FILL)

    ws4 = wb.create_sheet("Removed")
    r = _headers(ws4, 1, ["Room", "Family", "Area (sq.m)"], [34, 20, 18])
    _rows(ws4, r, [[x["name"], x["family"], x["old_sqm"]]
                   for x in cmp_data["removed"]], fill_when=lambda _r: WARN_FILL)

    wb.save(out_path)


WALL_LAYER = "AI_WALL"
HANGING_WALL_LAYER = "AI_HANGING_WALL"
# the per-class keys of a Generate saved before walls were one layer
_OLD_WALL_KEYS = ("external", "core", "internal")


def walls_summary(stats: Dict) -> Dict:
    """{AI_WALL: {"count", "length_m", "area_sqm"}} from a Generate's stats
    (walls_count / _length_m / _area_sqm, summed over floors). Stats saved
    before walls were one layer (walls_external_* / _core_* / _internal_*)
    are added up."""
    def total(what):
        v = stats.get("walls_%s" % what)
        if v is None:
            v = sum(float(stats.get("walls_%s_%s" % (k, what), 0) or 0) for k in _OLD_WALL_KEYS)
        return float(v or 0)
    return {WALL_LAYER: {"count": int(total("count")),
                         "length_m": round(total("length_m"), 1),
                         "area_sqm": round(total("area_sqm"), 2)}}


def hanging_walls_summary(stats: Dict) -> Dict:
    """The hanging walls of a Generate (walls drawn above floor level, an
    object under them): not walls, their area counted in the rooms.
    {"count", "area_sqm", "to_rooms_sqm", "items": [{floor, area_sqm,
    evidence, pts}]}; {} when there are none."""
    n = int(stats.get("hanging_walls_count") or 0)
    if not n:
        return {}
    return {"count": n,
            "area_sqm": round(float(stats.get("hanging_walls_area_sqm") or 0), 2),
            "to_rooms_sqm": round(float(stats.get("hanging_walls_to_rooms_sqm") or 0), 2),
            "items": stats.get("hanging_wall_items") or []}


def wall_trace_summary(stats: Dict) -> Dict:
    """The wall tracer's checks (walltrace.py), summed over floors."""
    if not stats.get("wall_trace_floors"):
        return {}
    pct = stats.get("wall_trace_hatched_pct_by_floor") or {}
    vals = [v for v in pct.values() if v is not None]
    return {"floors": int(stats.get("wall_trace_floors") or 0),
            "footprint_sqm": round(float(stats.get("wall_trace_footprint_sqm") or 0), 2),
            "run_m": round(float(stats.get("wall_trace_run_m") or 0), 1),
            "hatched_traced_pct": (round(sum(vals) / len(vals), 1) if vals else None),
            "hatched_traced_pct_by_floor": pct,
            "squeezed_kept": int(stats.get("wall_trace_strips_wall_squeezed") or 0),
            "ledges": int(stats.get("wall_trace_strips_ledge") or 0),
            "gaps": int(stats.get("wall_trace_strips_gap") or 0),
            "review": int(stats.get("wall_trace_strips_review") or 0) + int(stats.get("wall_trace_review_blocks") or 0),
            "shafts": int(stats.get("wall_trace_shafts") or 0),
            "columns": int(stats.get("wall_trace_columns") or 0),
            "review_items": stats.get("wall_review_items") or []}


def add_walls_sheet(path: str, summary: Dict, trace: Dict = None, hanging: Dict = None):
    """Append a Walls sheet to an existing workbook: the walls drawn on
    AI_WALL - pieces, centre-line length and area. Walls are never rooms
    and never in the areas. `hanging` (hanging_walls_summary): the walls
    drawn above floor level, on AI_HANGING_WALL - not walls, their area
    counted in the rooms."""
    from openpyxl import load_workbook
    if not summary:
        return
    wb = load_workbook(path)
    ws = wb.create_sheet("Walls")
    row = _title(ws, 1, "Walls drawn (never rooms, never counted in areas)")
    row = _headers(ws, row, ["Layer", "Pieces", "Length (m)", "Area (sq.m)"], [24, 10, 14, 14])
    data = [(k, v["count"], v["length_m"], v["area_sqm"]) for k, v in summary.items()]
    row = _rows(ws, row, data)
    if hanging:
        row = _title(ws, (row or ws.max_row) + 2,
                     "Hanging walls (an object drawn under them): not walls, counted in the rooms")
        row = _headers(ws, row, ["Layer", "Pieces", "Into rooms (sq.m)", "Area (sq.m)"], [24, 10, 14, 14])
        row = _rows(ws, row, [(HANGING_WALL_LAYER, hanging.get("count"), hanging.get("to_rooms_sqm"),
                               hanging.get("area_sqm"))])
        items = hanging.get("items") or []
        if items:
            row = _headers(ws, row, ["Floor", "Evidence", "From (x, y)", "To (x, y)"], [24, 10, 22, 22])
            data3 = []
            for it in items:
                xs = [p[0] for p in it["pts"]]; ys = [p[1] for p in it["pts"]]
                data3.append((it["floor"], it.get("evidence", ""), "%d, %d" % (min(xs), min(ys)),
                              "%d, %d" % (max(xs), max(ys))))
            row = _rows(ws, row, data3)
    if trace:
        row = _title(ws, (row or ws.max_row) + 2, "How the walls were decided (paired faces, confirmed by the wall hatch)")
        row = _headers(ws, row, ["Check", "Value", "", ""], [24, 10, 14, 14])
        checks = [("Floors traced", trace.get("floors")), ("Wall area (sq.m)", trace.get("footprint_sqm")),
                  ("Wall run (m)", trace.get("run_m")),
                  ("Wall hatch traced (%)", trace.get("hatched_traced_pct")),
                  ("Unhatched strips between two walls, kept", trace.get("squeezed_kept")),
                  ("Ledge walls behind WCs, left out", trace.get("ledges")),
                  ("Empty gaps, left out", trace.get("gaps")),
                  ("Hatched blocks taken as shafts", trace.get("shafts")),
                  ("Hatched blocks taken as columns", trace.get("columns")),
                  ("Pieces to review", trace.get("review"))]
        row = _rows(ws, row, [(a, b if b is not None else "-", "", "") for a, b in checks])
        items = trace.get("review_items") or []
        if items:
            row = _title(ws, (row or ws.max_row) + 2, "Wall pieces to review (mm)")
            row = _headers(ws, row, ["Floor", "Kind", "From (x, y)", "To (x, y)"], [24, 10, 22, 22])
            data2 = []
            for it in items:
                xs = [p[0] for p in it["pts"]]; ys = [p[1] for p in it["pts"]]
                data2.append((it["floor"], it["kind"], "%d, %d" % (min(xs), min(ys)), "%d, %d" % (max(xs), max(ys))))
            _rows(ws, row, data2)
    wb.save(path)


def add_wall_rules_sheet(path: str, stats: Dict):
    """Append a "Wall rules" sheet: every line the geometry rules took out
    of the walls (break mark, oblique unsupported, dead end) with its
    coordinates, the rooms grown to their stated size, and how many
    outlines were cleaned of slits and spikes."""
    from openpyxl import load_workbook
    exc = stats.get("wall_exclusions") or []
    grown = stats.get("stated_grown") or []
    wb = load_workbook(path)
    ws = wb.create_sheet("Wall rules")
    row = _title(ws, 1, "Lines that are not walls (by their geometry)")
    counts = {}
    for e in exc:
        counts[e["rule"]] = counts.get(e["rule"], 0) + 1
    for k in ("break mark", "oblique unsupported", "dead end"):
        row = _kv(ws, row, k, counts.get(k, 0))
    row = _kv(ws, row, "Outlines cleaned (slits, spikes)", int(stats.get("outlines_cleaned", 0) or 0))
    row = _kv(ws, row, "Rooms grown to their stated size", len(grown))
    row += 1
    row = _headers(ws, row, ["Floor", "Rule", "x1", "y1", "x2", "y2", "Length (mm)"],
                   [22, 22, 12, 12, 12, 12, 14])
    data = []
    for e in exc:
        x1, y1, x2, y2 = e["line"][:4]
        data.append((e["floor"], e["rule"], x1, y1, x2, y2, round(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5)))
    row = _rows(ws, row, data)
    if grown:
        row += 1
        row = _headers(ws, row, ["Floor", "Room grown", "Stated", "Before (sq.m)", "After (sq.m)"])
        _rows(ws, row, [(g["floor"], g["label"], g["stated"], g["before_sqm"], g["after_sqm"]) for g in grown])
    wb.save(path)


def add_window_checks_sheet(path: str, stats: Dict):
    """Append a "Windows" sheet: how the X-marked openings in the walls
    were sorted (windows, shaft strips, cut-outs), and every window whose
    layer says shaft or cut-out while the drawing shows glazing."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(path)
    ws = wb.create_sheet("Windows")
    row = _title(ws, 1, "X-marked openings in the walls")
    for label, key in (("Windows", "windows"), ("  of them told by glazing lines", "windows_by_glazing"),
                       ("Shaft strips (no glazing, shaft layer)", "shafts_from_window_test"),
                       ("Cut-outs drawn on AI_CUTOUT", "cutouts_drawn"),
                       ("Layer / drawing conflicts", "window_layer_conflicts")):
        row = _kv(ws, row, label, int((stats or {}).get(key, 0) or 0))
    row += 1
    row = _headers(ws, row, ["Floor", "Outline (mm)", "Glazing lines", "Layers", "Check"], [18, 34, 14, 50, 44])
    for r in (stats or {}).get("window_conflicts") or []:
        vals = [r["floor"], "{0}-{2} x {1}-{3}".format(*r["bbox"]), r["glazing"], ", ".join(r["layers"]), r["note"]]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=row, column=i, value=v).font = Font(bold=True, color="C00000")
        row += 1
    wb.save(path)


def add_size_check_sheet(path: str, rows):
    """Append a "Size check" sheet: every boundary whose label states a
    size, its drawn area against the stated one; ratios outside the
    boundary.size_check_range are flagged."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(path)
    ws = wb.create_sheet("Size check")
    row = _title(ws, 1, "Boundaries with a stated size: drawn against stated")
    row = _headers(ws, row, ["Floor", "Label", "Family", "Layer", "Stated (mm)", "Stated (sq.m)",
                             "Drawn (sq.m)", "Drawn / stated", "Check"], [18, 34, 20, 18, 16, 14, 14, 14, 24])
    for r in rows or []:
        vals = [r["floor"], r["label"], r["family"], r["layer"], r["stated"], r["stated_sqm"],
                r["drawn_sqm"], r["ratio"], r["flag"] or "OK"]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=i, value=v)
            if r["flag"]:
                c.font = Font(bold=True, color="C00000")
        row += 1
    wb.save(path)
