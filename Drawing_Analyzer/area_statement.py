# -*- coding: utf-8 -*-
"""Area Statement: every measured room filed under one canonical category,
one row per floor, the client's columns composed from those categories.

How it works
------------
1. Every room of the area result is classified into ONE canonical category
   (CATEGORIES): toilets, lift_lobby, service_lobby, staircase, lift_shaft,
   dg, phe, ... The classification uses, in this order, what the drawing
   says about the room:
     * how it was found - an X-marked slab opening is a cut-out and is
       deducted; a walled shaft room without an X (a PHE shaft room with a
       door) is a room, not an opening, and is listed but never deducted;
     * its name (the label the engine gave it);
     * its family (the boundary layer it was drawn on);
     * its neighbours - a lobby touching a service lift and no passenger
       lift is the service lobby.
2. Per floor plate the categories are summed; the plate is repeated for
   every floor its title names ("TYPICAL FLOOR PLAN-(2nd, 3rd & 4th)" is
   three rows). Plates are grouped into towers by the TOWER heading on the
   sheet, else the file name.
3. The statement's columns come from a template (config.json ->
   area_statement.templates): each column is one canonical category or the
   sum of several ("AHU shafts" on one client's sheet is the AHU rooms plus
   the AHU shafts; "STAIRCASE / Lobby PRESSURE SHAFT" is both pressure
   shaft kinds). The default template shows every canonical category some
   floor uses. Derived columns are formulas in the workbook:
     total shafts = sum of the opening columns
     Net BUA      = gross - total shafts
     plinth       = Net BUA - staircase - lift shaft
     sale area    = plinth / sale_area_factor (0.70), sq.ft at x10.764.

Areas are the outline areas of the drawn boundaries (what a manual
take-off measures); an opening inside a room is deducted once, in its own
column, never twice. Free of ezdxf; openpyxl only.
"""

from __future__ import annotations

import re
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import CONFIG
from .levels import parse_levels, tower_from_name
from .labels import declared_area

# -- canonical categories --------------------------------------------------
# (key, group, header). Groups: area, common, vertical, opening, info,
# derived. Openings are deducted (total shafts); info columns are listed
# for the reader and never deducted.
CATEGORIES = [
    ("gross", "area", "GROSS AREA\n(Pline)"),
    ("slab", "area", "SLAB AREA\n(Pline)"),
    ("elevation", "area", "ELEVATION\nAREA"),
    ("toilets", "common", "TOILETS"),
    ("lift_lobby", "common", "LIFT LOBBY"),
    ("service_lobby", "common", "SERVICE LOBBY"),
    ("jump_lobby", "common", "JUMP LIFTS LOBBY"),
    ("entrance_lobby", "common", "ENTRANCE LOBBY"),
    ("ahu_room", "common", "AHU ROOM"),
    ("ele_room", "common", "ELE ROOM"),
    ("refuge", "common", "REFUGE (Provided)"),
    ("balcony", "common", "BALCONY"),
    ("machine_room", "common", "MACHINE ROOM"),
    ("fire_command", "common", "FIRE COMMAND CENTER"),
    ("service_room", "common", "SERVICE ROOM"),
    ("staircase", "vertical", "STAIRCASE"),
    ("lift_shaft", "vertical", "LIFT SHAFT"),
    ("dg", "opening", "DG EXHAUST SHAFT"),
    ("toilet_exhaust", "opening", "TOILET EXHAUST SHAFT"),
    ("phe", "opening", "PHE / PL SHAFT"),
    ("comm", "opening", "COMMUNICATION SHAFT"),
    ("lobby_press", "opening", "LOBBY PRESS. SHAFT"),
    ("stair_press", "opening", "STAIRCASE PRESS. SHAFT"),
    ("rwdp", "opening", "RWDP SHAFT"),
    ("fp", "opening", "FP SHAFT"),
    ("fhr", "opening", "FHR / FHC SHAFT"),
    ("fapa", "opening", "FAPA SHAFT"),
    ("chw", "opening", "CHW SHAFT"),
    ("ahu_shaft", "opening", "AHU SHAFT"),
    ("pantry_shaft", "opening", "PANTRY SHAFT"),
    ("repro", "opening", "REPRO"),
    ("bms", "opening", "BMS"),
    ("ten", "opening", "TENANT RISER (TEN)"),
    ("stp", "opening", "STP EXHAUST AIR SHAFT"),
    ("lcs", "opening", "LCS"),
    ("spare", "opening", "SPARE SHAFT"),
    ("vrv", "opening", "VRV SHAFT"),
    ("elect_shaft", "opening", "ELECT. / EARTHING SHAFT"),
    ("condenser", "opening", "CONDENSER"),
    ("elv", "opening", "ELV SHAFT"),
    ("m_hvac", "opening", "M. HVAC"),
    ("future_shaft", "opening", "SHAFTS FOR FUTURE"),
    ("double_cutout", "opening", "DOUBLE HEIGHT CUTOUT"),
    ("triple_cutout", "opening", "TRIPLE HEIGHT CUTOUT"),
    ("void", "opening", "VOID"),
    ("other_opening", "opening", "OTHER NAMED SHAFTS"),
    ("unnamed_opening", "opening", "UNNAMED CUT-OUTS"),
    ("shaft_room", "info", "SHAFT ROOMS\n(walled, not deducted)"),
    ("wall_cutout", "info", "CUT-OUTS IN WALLS\n(not deducted)"),
    ("cutout_above", "info", "CUT-OUTS ABOVE / AT CEILING\n(not deducted)"),
    ("circulation", "info", "CORRIDORS / PASSAGES\n(not deducted)"),
    ("total_shafts", "derived", "TOTAL SHAFTS\n& CUTOUTS"),
    ("net_bua", "derived", "Net BUA"),
    ("plinth", "derived", "PLINTH AREA"),
    ("sale_sqm", "derived", "SALEABLE AREA\n(sq.m)"),
    ("sale_sft", "derived", "SALEABLE AREA\n(sq.ft)"),
    ("eff_sa", "derived", "BUA TO SA\nEFFICIENCY"),
    ("eff_pa", "derived", "GROSS TO PA\nEFFICIENCY"),
]
KEYS = [c[0] for c in CATEGORIES]
GROUP_OF = {c[0]: c[1] for c in CATEGORIES}
HEADER_OF = {c[0]: c[2] for c in CATEGORIES}
OPENING_KEYS = [k for k, g, _ in CATEGORIES if g == "opening"]
COMMON_KEYS = [k for k, g, _ in CATEGORIES if g == "common"]
VERTICAL_KEYS = ["staircase", "lift_shaft"]
# An X-marked opening near a walled shaft compartment but not along its
# outline (less than this, mm, within 60 mm) - behind its wall - and along
# another shaft compartment's outline is that other shaft's opening (XA:
# the ELE SH strips beside ELECTRICAL SH).
AGAINST_MIN_MM = 100.0
INFO_KEYS = [k for k, g, _ in CATEGORIES if g == "info"]
MEASURED = COMMON_KEYS + VERTICAL_KEYS + OPENING_KEYS + INFO_KEYS
DERIVED = ["total_shafts", "net_bua", "plinth", "sale_sqm", "sale_sft", "eff_sa", "eff_pa"]
# kept for older callers
SHAFT_KEYS = OPENING_KEYS
FIRST_COL = 2

# -- name patterns -----------------------------------------------------------
# Tried in order on the normalised name (upper case, punctuation to spaces).
# First match wins, so the specific comes before the general.
SHAFT_PATTERNS: List[Tuple[str, List[str]]] = [
    ("double_cutout", [r"DOUBLE HEIGHT"]),
    ("triple_cutout", [r"TRIPLE HEIGHT"]),
    ("lobby_press", [r"\bLOBBY\b.*\bPRE[SZ]", r"\bLOBBY\s+PR\b", r"\bLIFT\b.*\bPRE[SZ]?\w*\b", r"\bLLP\b"]),
    ("stair_press", [r"\bSTAIR\w*\b.*\bPR", r"\bPRESS\w*\b.*\bSTAIR", r"^PRESS\w*$", r"\bPRESS\w*\s*SH", r"\bSHAFT\s*PRESS", r"\bPREZZ?\b", r"\bSTR\s*PRE\b"]),
    ("toilet_exhaust", [r"\bTOI\w*\b.*\bEX", r"\bT\.?\s*EX\b", r"\bTE SHAFT\b",
                        r"\bSHAFT\s+FOR\s+TOI"]),     # H10 T2: SHAFT FOR TOILET 1200X1175
    ("dg", [r"\bDG\b"]),
    ("stp", [r"\bSTP\b"]),
    ("fapa", [r"\bFA\s*PA\b", r"\bPA\s*FA\b"]),
    ("fhr", [r"\bFHR\b", r"\bFHC\b", r"\bHOSE\s*REEL"]),
    ("fp", [r"\bFP\b", r"\bFIRE\s*PUMP\b"]),
    ("chw", [r"\bCHW\b", r"\bCHILL\w*\b", r"\bCH\s*PIPE\b", r"\bTFA\b"]),
    ("rwdp", [r"\bRWDP\b", r"\bRW\w*\b.*\bSHAFT", r"\bRWP\b", r"\bRAIN\b"]),
    ("phe", [r"\bPHE\b", r"\bPL\s*SHAFT\b", r"^PL\W*SH\b", r"\bPLUMB", r"\bPLB\b"]),
    ("comm", [r"\bCOMM\w*\b", r"^CO\s*SH\b"]),
    ("bms", [r"\bBMS\b"]),
    ("ten", [r"\bTEN\b", r"\bTENANT\b", r"^TR\s*SH\b"]),
    ("lcs", [r"\bLCS\b"]),
    ("elv", [r"\bELV\b"]),
    ("vrv", [r"\bVRV\b"]),
    ("repro", [r"\bREPRO\b"]),
    ("spare", [r"\bSPARE\b"]),
    ("future_shaft", [r"\bFUTURE\b"]),
    ("condenser", [r"\bCONDENS"]),
    ("ahu_shaft", [r"\bAHU\b.*\b(SHAFT|PIPE|DUCT)\b", r"\bAHU\s*S\b"]),
    ("pantry_shaft", [r"\bPANTRY\b.*\b(SHAFT|DUCT|EX)"]),
    ("elect_shaft", [r"^EL\W*ELE\b",                 # H10 T2: EL-ELE-03 1200X1175
                     r"\bELE\w*\b.*\bSHAFT", r"\bELE\w*\s*SH\b", r"^EL\s*SH\b", r"\bE\s*SHAFT\b", r"\bEARTH", r"^ELEC?$", r"\bELECT\w*\s*CUT"]),
    ("m_hvac", [r"\bHVAC\b"]),
    ("void", [r"^VOID\b"]),
]
ROOM_PATTERNS: List[Tuple[str, List[str]]] = [
    ("fire_command", [r"\bFIRE\s*COMMAND\b", r"\bFCC\b"]),
    ("machine_room", [r"\bMACHINE\s*ROOM\b", r"\bLMR\b", r"\bLIFT\s*MACHINE\b",
                      r"^M\s*\.?\s*ROOM\b"]),          # H10 T2 terrace: M. ROOM 3000X4150
    ("service_lobby", [r"\bSERVICE\s*LOBB", r"\bFIRE\s*TOWER\b", r"\bFIRE\w*\s*LOBB", r"\bFIRE\s*LIFT\s*LOBB"]),
    ("jump_lobby", [r"\bJUMP\b.*\bLOBB"]),
    ("entrance_lobby", [r"\bENTRANCE\b.*\bLOBB", r"\bGROUND\s*LOBB"]),
    ("ahu_room", [r"\bAHU\b", r"\bAIR\s*HANDL"]),
    ("ele_room", [r"\bELE\w*\s*ROOM\b", r"\bELE\.?\s*ROOM", r"\bELEC\w*\b.*\bROOM\b", r"\bPANEL\s*ROOM\b", r"\bLT\s*ROOM\b"]),
    ("refuge", [r"\bREFUGE\b"]),
    ("balcony", [r"\bBALCON", r"\bTERRACE\b", r"\bDECK\b"]),
    ("toilets", [r"\bTOILET", r"\bTOI\b", r"\bJANIT", r"\bJAN\b", r"\bPCT\b", r"\bP\.?C\b.*\bTOI", r"\bWASH\s*ROOM", r"\bREST\s*ROOM", r"\bWC\b", r"\bLADIES\b", r"\bGENTS\b", r"\bURINAL", r"\bHANDICAP", r"\bH\.?\s*TOILET"]),
    ("staircase", [r"\bSTAIR", r"\bST\s*-?\s*\d+\b", r"\bFIRE\s*ESCAPE\b"]),
    ("lift_lobby", [r"\bLIFT\s*LOBB", r"\bELEVATOR\s*LOBB"]),
    ("lift_shaft", [r"\bLIFT\b", r"\bELEVATOR\b", r"\bPAX\b"]),
    ("service_room", [r"\bSERVICE\s*ROOM\b"]),
    ("circulation", [r"\bLOBB", r"\bCORRIDOR", r"\bPASSAGE", r"\bFOYER", r"\bCIRCULATION"]),
]
# Names an X-marked opening gets from a tag standing in it: a column mark
# (C-7), a door tag (SD-6), a block-work note - never a shaft's name.
TAG_NAME = re.compile(r"^(C|SD|D|W|FMD|R|S|G)\s*-?\s*\d+[A-Z]?$|^UNKNOWN|CUT\s*OUT|DROP|SLEEVE|BLOCK\s*WORK|^\d+\s*(X|MM)|^-?$")

FAMILY_CATEGORY = {
    "Toilet": "toilets", "Staircase": "staircase", "Lift": "lift_shaft",
    "Lift Lobby": "lift_lobby", "Service Lobby": "service_lobby",
    "Electrical": "ele_room", "Balcony / Open": "balcony", "Refuge": "refuge",
    "Circulation": "circulation", "AHU": "ahu_room", "AHU Room": "ahu_room",
    # a lift machine room (LIFT ELE.PANELS), and a lift well with no door
    # on it (dxf_reader rule lift_no_door_machine_room): MACHINE ROOM, not
    # usable area and not LIFT SHAFT (Test_2, decision 3b)
    "Machine Room": "machine_room",
    # a lift with no door opening (dxf_reader rule lift_without_door_mechanical)
    "Mechanical Room": "machine_room",
}
LIFT_LAYER_WORDS = ("LIFT",)


def _norm(name: str) -> str:
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", (name or "").upper()).split())


def _first(patterns, n):
    for key, pats in patterns:
        if any(re.search(p, n) for p in pats):
            return key
    return None


def _settings(cfg=None):
    s = (cfg or CONFIG).get("area_statement") or {}
    return {
        "sale_factor": float(s.get("sale_area_factor", 0.70)),
        "sqft": float(s.get("sqft_per_sqm", 10.764)),
        "deduct_unnamed": bool(s.get("deduct_unnamed_openings", True)),
        "templates": s.get("templates") or {},
        "template": s.get("template") or "default",
        "min_opening_sqm": float(s.get("min_unnamed_opening_sqm", 0.15)),
        "use_stated_area": bool(s.get("use_stated_area", True)),
        "stated_area_tol": float(s.get("stated_area_tolerance", 0.03)),
    }


def classify(name: str, family: str = "", opening=None,
             neighbours=None, within_family: str = "", area_sqm: float = 0.0) -> Optional[str]:
    """The canonical category of a room, or None for usable area.

    opening: True when the boundary is an X-marked slab opening (dxf_reader
      tags every boundary it draws), False when it was traced from walls,
      None when the drawing does not say (a boundary drawn by hand or by an
      older version): then a shaft by name or family counts as an opening,
      as it always did.
    neighbours: [{"name", "layer"}] of the rooms touching this one.
    within_family: the family of the room this one stands inside, if any.
    """
    n = _norm(name)
    fam = family or ""
    if fam == "Void":
        # a void - the X sequence's (every X mark decided once, 2026-10-16)
        # or drawn so: filed as a void whatever the text in it names (a
        # landscape note, a terrace): a double / triple height cut-out by
        # its name, a cut-out in the slab above as one; the well of a stair
        # or a lift is part of it
        if within_family in ("Staircase", "Lift"):
            return None
        if re.search(r"\bABOVE\b|\bAT\s+CEILING\b", n) and re.search(r"\bCUT|\bCEILING\b|\bOPENING\b", n):
            return "cutout_above"
        key = _first(SHAFT_PATTERNS, n)
        return key if key in ("double_cutout", "triple_cutout", "void") else "void"
    if opening is None:
        opening = bool(fam in ("Shaft", "Void", "Cutout") or _first(SHAFT_PATTERNS, n)
                       or re.search(r"\bSHAFT\b|\bCUT\s*OUT\b|\bRISER\b", n))
    # 1. a slab opening: deducted under its name, or as an unnamed cut-out
    if opening:
        if within_family in ("Staircase", "Lift"):
            return None                    # the well of a stair / lift: part of it
        if fam in ("Machine Room", "Mechanical Room"):
            return "machine_room"          # a lift well with no door, part of the machine room (P285 11th floor)
        if re.search(r"\bABOVE\b|\bAT\s+CEILING\b", n) and re.search(r"\bCUT|\bCEILING\b|\bOPENING\b", n):
            return "cutout_above"          # a cut-out in the slab above (BHUB: CUTOUT ABOVE AT CEILING), not in this floor
        key = _first(SHAFT_PATTERNS, n)
        if key:
            return key
        if re.search(r"\bLIFT\b|\bPAX\b|\bELEVATOR\b", n) and fam in ("Lift", ""):
            return "lift_shaft"
        if fam == "Lift":
            return "lift_shaft"
        if fam == "Staircase" or re.search(r"\bSTAIR", n):
            return "staircase"
        if fam == "Toilet" or re.search(r"\bTOILET|\bWC\b", n):
            return "toilets"
        if _first([p_ for p_ in ROOM_PATTERNS if p_[0] == "machine_room"], n):
            return "machine_room"          # a lift machine room over its X'd well (XA tower, 22nd floor)
        if TAG_NAME.search(n) or not n or re.search(r"\bSHAFT\b|\bCUT", n):
            return "unnamed_opening"
        if fam in ("Shaft", "MEP / Services", "AHU Room", "Plumbing", "Fire / Life Safety", "Unknown", ""):
            return "other_opening"
        return "unnamed_opening"
    # 2. a walled shaft compartment (no X of its own). With X-marked
    #    openings inside or against it, the openings are what is deducted
    #    (a PHE shaft room with two slab openings); with none drawn, the
    #    compartment is the shaft (an AHU pipe shaft's kerb) and is deducted.
    if fam in ("Shaft", "MEP / Services", "Plumbing", "Fire / Life Safety", "Unknown", "", "Cutout", "Void"):
        key = _first(SHAFT_PATTERNS, n)
        if key == "fhr" and re.search(r"\bFHC\b", n) and not re.search(r"\bSHAFT|\bRISER|\bDUCT", n):
            # a hose cabinet (FHC-3) is not a shaft; its X'd slot is - and
            # the cabinet in front of the slot, behind its door, is the
            # slot's room (Test_2)
            if any(x.get("opening") and _shared_near(x) >= AGAINST_MIN_MM and
                   _first(SHAFT_PATTERNS, _norm(x.get("name"))) == "fhr" for x in neighbours or []):
                return "shaft_room"
            key = None
        if key or re.search(r"\bSHAFT\b|\bRISER\b|\bDUCT\b", n):
            # the openings of this shaft: inside it, or against it and not
            # another shaft's (unnamed, or named for the same category)
            for x in neighbours or []:
                if not x.get("opening"):
                    continue
                if x.get("layer") in ("AI_LIFT", "AI_STAIR", "AI_VOID") or (
                        re.search(r"\bLIFT\b|\bPAX\b|\bELEVATOR\b", _norm(x.get("name")))
                        and not _first(SHAFT_PATTERNS, _norm(x.get("name")))):
                    continue                # a lift well's X (found by its shape) is not this shaft's opening
                if x.get("inside"):
                    return "shaft_room"
                if float(x.get("shared_mm") or 0.0) < AGAINST_MIN_MM and any(
                        _first(SHAFT_PATTERNS, _norm(t)) or re.search(r"\bSHAFT\b|\bRISER\b|\bDUCT\b", _norm(t))
                        for t in (x.get("touches") or ())):
                    continue                # behind this one's wall, along another shaft's: that one's opening
                xn = _norm(x.get("name"))
                xkey = _first(SHAFT_PATTERNS, xn)
                if xkey is None or xkey == key or TAG_NAME.search(xn) or not xn:
                    return "shaft_room"
            return key or "other_opening"
    key = _first(ROOM_PATTERNS, n)
    if key == "lift_shaft" and fam not in ("Lift", "", "Unknown"):
        key = None                          # "LIFT LOBBY PRE..." handled above; a Lift-named room of another family
    if key == "circulation":
        key = _lobby_kind(n, neighbours)
    if key == "staircase" and fam not in ("Staircase", "", "Unknown"):
        key = FAMILY_CATEGORY.get(fam, key)
    if key:
        return key
    # 3. by family alone
    key = FAMILY_CATEGORY.get(fam)
    if key == "circulation":
        return _lobby_kind(n, neighbours)
    if key == "lift_shaft":
        return "lift_shaft"
    return key


def _gross_of(floor) -> float:
    """The floor's gross area: its outline, or - when the outline is the
    slab edge (a slab-projection layer) and the facade could be traced -
    the area to the outer face of the facade (facade.py)."""
    g = floor.get("glass_face_area_sqm")
    if g and _setting("gross_to_facade", True):
        return float(g)
    return float(floor.get("area_sqm") or 0.0)


def _setting(key, default):
    try:
        from .config import CONFIG
        return (CONFIG.get("area_statement") or {}).get(key, default)
    except Exception:                                 # noqa: BLE001
        return default


def _shared_near(x) -> float:
    """How far a neighbour's outline runs along this one, allowing the two
    a small gap (dxf_reader.SHARED_NEAR_MM): the cabinet in front of its
    X'd slot, a thin panel between them (P285 FHC, shared_mm 0)."""
    return max(float(x.get("shared_mm") or 0.0), float(x.get("shared_near_mm") or 0.0))


def _owner_shaft(r) -> Optional[str]:
    """The shaft category of the named, non-opening shaft compartment this
    unnamed opening stands in, or lies along for most of its long side (a
    strip across the end of a PHE shaft), if there is exactly one. An
    electrical room counts as the electrical shaft's compartment."""
    a = float(r.get("gross_area_sqm") or r.get("area_sqm") or 0.0) * 1e6
    half = float(r.get("perimeter_m") or 0.0) * 1000.0 / 2.0
    disc = max(half * half - 4.0 * a, 0.0) ** 0.5
    long_side = (half + disc) / 2.0 if half else 0.0
    found = set()
    for x in r.get("neighbours") or []:
        if x.get("opening"):
            continue
        key = _first(SHAFT_PATTERNS, _norm(x.get("name")))
        if not key and _first(ROOM_PATTERNS, _norm(x.get("name"))) == "ele_room":
            key = "elect_shaft"             # the riser slot along an electrical room (BHUB)
        if not key or key in ("void", "double_cutout", "triple_cutout"):
            continue
        shared = float(x.get("shared_mm") or 0.0)
        if long_side and shared < 0.8 * long_side:
            continue
        found.add(key)
    return next(iter(found)) if len(found) == 1 else None


def _lobby_kind(n: str, neighbours) -> str:
    """A lobby's category from what it touches: service lifts only ->
    service lobby; passenger lifts -> lift lobby; nothing -> circulation."""
    if re.search(r"\bSERVICE\b|\bFIRE\b", n):
        return "service_lobby"
    lifts = [x for x in (neighbours or []) if "LIFT" in (x.get("layer") or "").upper()
             or re.search(r"\bLIFT\b|\bPAX\b|\bELEVATOR\b", _norm(x.get("name")))]
    lifts = [x for x in lifts if not re.search(r"\bPRE(S|Z|\b)|\bLOBB", _norm(x.get("name")))]
    if not lifts or not re.search(r"\bLOBB|\bLIFT\b", n):
        return "lift_lobby" if re.search(r"\bLIFT\b", n) else "circulation"
    service = [x for x in lifts if re.search(r"\bSERVICE\b|\bFIRE\b|\bGOODS\b|\bFREIGHT\b", _norm(x.get("name")))]
    if service and len(service) == len(lifts):
        return "service_lobby"
    return "lift_lobby"


# -- floors and towers -------------------------------------------------------
_LEVEL_PATTERNS = [
    (r"\b(\d{1,3})\s*(ST|ND|RD|TH)\s+FLOOR\b",
     lambda m: "{0}{1} FLOOR".format(m.group(1), m.group(2).lower())),
    (r"\bGROUND\s+FLOOR\b", lambda m: "GROUND FLOOR"),
    (r"\bBASEMENT\s*[-_]?\s*(\d+)\b", lambda m: "BASEMENT {0}".format(m.group(1))),
    (r"\bTERRACE\b", lambda m: "TERRACE"),
    (r"\bLEVEL\s*[-_]?\s*(\d+)\b", lambda m: "LEVEL-{0}".format(m.group(1))),
]


def level_name(drawing_name: str) -> str:
    """The floor a drawing shows, from its file name; the file name when
    it names none."""
    raw = re.sub(r"\.dxf$", "", drawing_name or "", flags=re.I)
    text = re.sub(r"[_]+", " ", raw).upper()
    for pat, fmt in _LEVEL_PATTERNS:
        m = re.search(pat, text)
        if m:
            return fmt(m)
    return raw.strip() or "FLOOR"


def _plate_levels(f, drawing_name):
    levels = list(f.get("levels") or [])
    if levels:
        return levels
    levels = parse_levels(f.get("title") or "") or parse_levels(level_name(drawing_name))
    if levels:
        return levels
    return [{"key": 0, "label": str(f.get("name") or "Floor"), "refuge": False}]


def _room_area(r) -> float:
    """The outline area: what a take-off polyline measures. Openings inside
    it are deducted in their own column."""
    g = r.get("gross_area_sqm")
    return float(g if g is not None else (r.get("area_sqm") or 0.0))


def build_statement(area_data: Dict, drawing_name: str = "", cfg=None,
                    template: Optional[str] = None) -> Dict:
    """Towers, each with one row per floor, plus the room list.

    Each row carries "canon": every canonical category summed for that
    plate, "values": the template's columns, "gross": the plate area.
    """
    st = _settings(cfg)
    floors = area_data.get("floors") or []
    if not floors:
        floors = [{"name": area_data.get("floor_label") or "Floor",
                   "area_sqm": area_data.get("floor_area_sqm") or 0.0}]
    file_tower = tower_from_name(drawing_name)
    tpl = _template(st, template)

    plates = []
    for f in floors:
        sums = {k: 0.0 for k in MEASURED}
        rooms = []
        for r in area_data.get("rooms") or []:
            if len(floors) > 1 and r.get("floor") != f.get("name"):
                continue
            key = classify(r.get("name", ""), r.get("family", ""),
                           opening=(bool(r.get("opening")) if r.get("source") else None),
                           neighbours=r.get("neighbours"),
                           within_family=r.get("within_family") or "",
                           area_sqm=float(r.get("area_sqm") or 0.0))
            a = _room_area(r)
            note = ""
            # The label states the room's area ("REFUGE AREA 151.05 SQ.
            # MTR."): that is the architect's figure, and the one the
            # statement carries; the traced outline is shown beside it.
            # H10 Tower 2's refuge outlines run open along the facade.
            stated = declared_area(r.get("name", "")) if (key and st["use_stated_area"]) else None
            if stated and a > 0 and abs(a - stated) > st["stated_area_tol"] * stated:
                note = "label states {0:.2f} sq.m; traced {1:.2f}".format(stated, a)
                a = stated
            if key == "unnamed_opening":
                # an opening against a named shaft compartment, or inside
                # it, is that shaft's opening: filed under its name (the
                # PHE shaft on 285 FD is one label, two X-marked strips)
                owner = _owner_shaft(r)
                if owner:
                    key, note = owner, "opening of the shaft beside it"
            if key == "unnamed_opening" and not st["deduct_unnamed"]:
                key, note = None, "unnamed opening, not deducted (config)"
            if key == "unnamed_opening" and a < st["min_opening_sqm"]:
                key, note = None, "a sleeve, below min_unnamed_opening_sqm"
            if key:
                sums[key] += a
            rooms.append((r, key, a, note))
        # Cut-outs in the walls (AI_CUTOUT) are openings in walls, not in
        # the slab: listed, never deducted - unless named as a double /
        # triple height cut-out, which is a slab opening.
        for c in area_data.get("cutouts") or []:
            if len(floors) > 1 and c.get("floor") != f.get("name"):
                continue
            nm = _norm(c.get("name") or "")
            key = _first(SHAFT_PATTERNS, nm)
            if key not in ("double_cutout", "triple_cutout"):
                key = "wall_cutout"
            a = float(c.get("area_sqm") or 0.0)
            sums[key] += a
            rooms.append(({"name": c.get("name") or "CUTOUT", "family": "Cutout",
                           "area_sqm": a}, key, a, "cut-out in a wall, not deducted" if key == "wall_cutout" else "wall cut-out"))
        canon = {k: round(v, 3) for k, v in sums.items()}
        canon["gross"] = _gross_of(f)
        canon["slab"] = float(f.get("area_sqm") or 0.0)
        plates.append({"floor": f, "levels": _plate_levels(f, drawing_name),
                       "canon": canon, "values": _compose(canon, tpl),
                       "rooms": rooms, "tower": f.get("tower") or file_tower})

    # towers: a plate whose floors are already taken in its tower is another tower
    towers = []
    for p in plates:
        keys = {lv["key"] for lv in p["levels"]}
        base = p["tower"] or "TOWER"
        for t in towers:
            if t["base"] == base and not (t["keys"] & keys):
                t["plates"].append(p)
                t["keys"] |= keys
                break
        else:
            towers.append({"base": base, "plates": [p], "keys": set(keys)})
    counts, seen = {}, {}
    for t in towers:
        counts[t["base"]] = counts.get(t["base"], 0) + 1
    for t in towers:
        seen[t["base"]] = seen.get(t["base"], 0) + 1
        if t["base"] == "TOWER":
            t["title"] = "TOWER {0}".format(seen["TOWER"]) if counts["TOWER"] > 1 else "TOWER 1"
        elif counts[t["base"]] > 1:
            t["title"] = "{0} ({1})".format(t["base"], seen[t["base"]])
        else:
            t["title"] = t["base"]

    out_towers, rooms_out = [], []
    for t in towers:
        rows = []
        for p in t["plates"]:
            for lv in p["levels"]:
                label = lv["label"] + (" (Refuge)" if lv.get("refuge") else "")
                rows.append({"key": lv["key"], "level": label,
                             "gross": _gross_of(p["floor"]),
                             "values": dict(p["values"]), "canon": dict(p["canon"]),
                             "plate": p["floor"].get("title") or p["floor"].get("name", "")})
            floors_txt = ", ".join(lv["label"] for lv in p["levels"])
            for r, key, a, note in p["rooms"]:
                col = tpl["column_of"].get(key) if key else None
                rooms_out.append({"tower": t["title"], "level": floors_txt,
                                  "plate": p["floor"].get("title") or p["floor"].get("name", ""),
                                  "layer": r.get("layer") or "",
                                  "name": r.get("name", ""), "family": r.get("family", ""),
                                  "area_sqm": round(a, 2), "key": key,
                                  "deducted": bool(key in OPENING_KEYS),
                                  "column": (tpl["header_of"].get(col) if col else
                                             (HEADER_OF.get(key) if key else None)),
                                  "note": note or ("" if key else "usable area"),
                                  "confidence": r.get("confidence"),
                                  "confidence_band": r.get("confidence_band") or "",
                                  "confidence_why": r.get("confidence_why") or ""})
        rows.sort(key=lambda r: r["key"])
        out_towers.append({"title": t["title"], "rows": rows})
    out_towers.sort(key=lambda t: _natural(t["title"]))
    return {"towers": out_towers, "rooms": rooms_out, "template": tpl,
            "sale_factor": st["sale_factor"], "sqft": st["sqft"]}


def _natural(text):
    return [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", text)]


# -- templates ---------------------------------------------------------------
def _template(st, name=None):
    """{"columns": [{"key", "header", "group", "sum": [canon keys]}],
    "column_of": canon key -> column key, "header_of": column key -> header}.

    A template in config names its columns; "sum" lists the canonical
    categories a column adds up (one by default). Categories no column
    names fall into their own column (so nothing measured is lost)."""
    name = name or st["template"]
    spec = (st["templates"] or {}).get(name) or {}
    cols = []
    for c in spec.get("columns") or []:
        key = c.get("key") or (c.get("sum") or [None])[0]
        if not key:
            continue
        cols.append({"key": key, "header": c.get("header") or HEADER_OF.get(key, key),
                     "group": c.get("group") or GROUP_OF.get(key, "common"),
                     "sum": list(c.get("sum") or [key])})
    covered = {k for c in cols for k in c["sum"]}
    for key in MEASURED:
        if key not in covered:
            cols.append({"key": key, "header": HEADER_OF[key], "group": GROUP_OF[key], "sum": [key]})
    order = {k: i for i, k in enumerate(KEYS)}
    cols.sort(key=lambda c: min(order.get(k, 999) for k in c["sum"]))
    column_of = {}
    for c in cols:
        for k in c["sum"]:
            column_of[k] = c["key"]
    # Columns a project also takes off the plinth (H10 Tower 3: refuge).
    less = [k for k in (spec.get("plinth_less") or []) if any(c["key"] == k for c in cols)]
    return {"name": name, "columns": cols, "column_of": column_of,
            "header_of": {c["key"]: c["header"] for c in cols},
            "group_of": {c["key"]: c["group"] for c in cols},
            "plinth_less": less}


def _compose(canon, tpl):
    return {c["key"]: round(sum(canon.get(k, 0.0) for k in c["sum"]), 2) for c in tpl["columns"]}


def derive(values: Dict, gross: float, tpl, sale_factor: float, sqft: float) -> Dict:
    """The derived figures for one row, from the template's columns."""
    shafts = sum(v for k, v in values.items() if tpl["group_of"].get(k) == "opening")
    vert = sum(v for k, v in values.items() if tpl["group_of"].get(k) == "vertical")
    vert += sum(values.get(k, 0.0) or 0.0 for k in tpl.get("plinth_less") or []
                if tpl["group_of"].get(k) != "vertical")
    net = gross - shafts
    plinth = net - vert
    sale = plinth / sale_factor if sale_factor else 0.0
    return {"total_shafts": shafts, "net_bua": net, "plinth": plinth,
            "sale_sqm": sale, "sale_sft": sale * sqft,
            "eff_sa": gross / sale if sale else 0.0,
            "eff_pa": gross / plinth if plinth else 0.0}


# -- workbook ----------------------------------------------------------------
NAVY = "1F3A5F"
_THIN = Side(border_style="thin", color="8EA0B8")
BORDER = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
GROUP_FILL = PatternFill("solid", fgColor="D9E2F1")
TITLE_FILL = PatternFill("solid", fgColor="E6EEF8")
WHITE_BOLD = Font(bold=True, color="FFFFFF", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
KEY_FILLS = {
    "gross": PatternFill("solid", fgColor="DDEBF7"),
    "total_shafts": PatternFill("solid", fgColor="FCE4E4"),
    "net_bua": PatternFill("solid", fgColor="E2EFDA"),
    "plinth": PatternFill("solid", fgColor="E2EFDA"),
    "sale_sqm": PatternFill("solid", fgColor="FFF2CC"),
    "sale_sft": PatternFill("solid", fgColor="FFF2CC"),
    "eff_sa": PatternFill("solid", fgColor="EDE7F6"),
    "eff_pa": PatternFill("solid", fgColor="EDE7F6"),
}
INFO_FILL = PatternFill("solid", fgColor="F2F2F2")
TOTAL_BAND = PatternFill("solid", fgColor="FFD966")
REFUGE_RED = "C00000"
GROUP_TITLE = {"common": "COMMON AREA", "vertical": "VERTICAL CONNECTIVITY",
               "opening": "SHAFTS & CUTOUTS", "info": "FOR INFORMATION (not deducted)",
               "derived": "BUA, PA, SA AND EFFICIENCY", "area": None}


def _active_columns(data):
    """The template columns some row uses, in template order, between the
    fixed head and tail."""
    used = set()
    for t in data["towers"]:
        for r in t["rows"]:
            used |= {k for k, v in r["values"].items() if abs(v) > 0.004}
    return [c for c in data["template"]["columns"] if c["key"] in used]


class _Sheet:
    def __init__(self, cols, tpl):
        self.cols = [{"key": "sl", "header": "Sl No", "group": None},
                     {"key": "floor", "header": "OFFICE FLOORS", "group": None},
                     {"key": "gross", "header": HEADER_OF["gross"], "group": None}] + \
                    [dict(c) for c in cols] + \
                    [{"key": k, "header": HEADER_OF[k], "group": "derived"} for k in DERIVED]
        self.keys = [c["key"] for c in self.cols]
        self.col = {k: get_column_letter(FIRST_COL + i) for i, k in enumerate(self.keys)}
        self.last = get_column_letter(FIRST_COL + len(self.keys) - 1)
        self.tpl = tpl

    def has(self, k):
        return k in self.col

    def ref(self, k, r):
        return "{0}{1}".format(self.col[k], r)

    def formulas(self, r, sale_factor, sqft):
        openings = [self.ref(c["key"], r) for c in self.cols if c.get("group") == "opening"]
        less = (self.tpl or {}).get("plinth_less") or []
        vert = "".join("-" + self.ref(c["key"], r) for c in self.cols
                       if c.get("group") == "vertical" or (c["key"] in less and self.has(c["key"])))
        g = self.ref("gross", r)
        return {
            "total_shafts": "=" + ("+".join(openings) if openings else "0"),
            "net_bua": "={0}-{1}".format(g, self.ref("total_shafts", r)),
            "plinth": "={0}{1}".format(self.ref("net_bua", r), vert),
            "sale_sqm": "={0}/{1}".format(self.ref("plinth", r), sale_factor),
            "sale_sft": "={0}*{1}".format(self.ref("sale_sqm", r), sqft),
            "eff_sa": "=IFERROR({0}/{1},0)".format(g, self.ref("sale_sqm", r)),
            "eff_pa": "=IFERROR({0}/{1},0)".format(g, self.ref("plinth", r)),
        }


def _band(ws, sh, row, text, size=12, fill=TITLE_FILL, height=None, font_color=NAVY):
    ws.merge_cells("{0}{1}:{2}{1}".format(get_column_letter(FIRST_COL), row, sh.last))
    c = ws.cell(row=row, column=FIRST_COL, value=text)
    c.font = Font(bold=True, size=size, color=font_color)
    c.alignment = CENTER
    c.fill = fill
    if height:
        ws.row_dimensions[row].height = height
    return row + 1


def _headers(ws, sh, row):
    groups = []
    for i, c in enumerate(sh.cols):
        col = FIRST_COL + i
        group = GROUP_TITLE.get(c.get("group")) if c.get("group") else None
        if group and groups and groups[-1][0] == group:
            groups[-1][2] = col
        elif group:
            groups.append([group, col, col])
        else:
            ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
        cell = ws.cell(row=row + (1 if group else 0), column=col, value=c["header"])
        cell.font = WHITE_BOLD
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
    for group, c0, c1 in groups:
        if c1 > c0:
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
        cell = ws.cell(row=row, column=c0, value=group)
        cell.font = Font(bold=True, color=NAVY)
        cell.fill = GROUP_FILL
        cell.alignment = CENTER
    for r in (row, row + 1):
        for i in range(len(sh.keys)):
            ws.cell(row=r, column=FIRST_COL + i).border = BORDER
    ws.row_dimensions[row + 1].height = 48
    return row + 2


def _style_row(ws, sh, r, refuge=False, total=False):
    for i, c in enumerate(sh.cols):
        k = c["key"]
        cell = ws.cell(row=r, column=FIRST_COL + i)
        cell.border = BORDER
        if total:
            cell.fill = TOTAL_BAND
        elif k in KEY_FILLS:
            cell.fill = KEY_FILLS[k]
        elif c.get("group") == "info":
            cell.fill = INFO_FILL
        if total:
            cell.font = Font(bold=True, color=NAVY)
        if refuge:
            cell.font = Font(bold=(k == "floor"), color=REFUGE_RED)
        if k in ("eff_sa", "eff_pa"):
            cell.number_format = "0.00"
        elif k not in ("sl", "floor"):
            cell.number_format = "#,##0.00"


def _family_layer_of(family: str, layer: str = "") -> str:
    """The AI_ layer of a room: the one it was drawn on, else its family's
    (family_rules.json)."""
    if layer:
        return layer
    try:
        from . import dxf_reader as _dr
        fam = _dr.family_by_name(family, _dr._FAMILY_RULES) if family else None
        if fam:
            return _dr._family_layer(fam)
    except Exception:                                 # noqa: BLE001
        pass
    return ""


INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
PLATE_FILL = PatternFill("solid", fgColor="BDD7EE")
COLUMN_FILL = PatternFill("solid", fgColor="F2F4F7")
LOOKUP_FILL = PatternFill("solid", fgColor="465A7A")
KEY_FONT = Font(italic=True, size=8, color="A0A0A0")
REG_KEY_FONT = Font(size=9, color="7F7F7F")
GROUP_KEYS = [("area", "GROSS AREA"), ("common", "COMMON AREA"), ("vertical", "VERTICAL CONNECTIVITY"),
              ("opening", "SHAFTS & CUTOUTS"), ("info", "FOR INFORMATION (not deducted)"),
              ("derived", "BUA, PA, SA AND EFFICIENCY"), ("usable", "USABLE (not in statement)")]
GROUP_NAME = {"area": "Area", "common": "Common", "vertical": "Vertical", "opening": "Opening",
              "info": "Info", "derived": "Derived", "usable": "Usable"}
MAP_SKIP = ("slab", "elevation")
FLOOR_SPARE_ROWS = 5


def _map_rows():
    """The category map: (key, header, group name, typical layer, deducted, plinth-less)."""
    typical = {"gross": "AI_FLOOR", "toilets": "AI_TOILET", "lift_lobby": "AI_LIFTLOBBY",
               "service_lobby": "AI_SERVICELOBBY", "jump_lobby": "AI_CIRC", "entrance_lobby": "AI_CIRC",
               "ahu_room": "AI_AHUROOM", "ele_room": "AI_ELEC", "refuge": "AI_FIRE", "balcony": "AI_BALCONY",
               "machine_room": "AI_MACHINEROOM", "fire_command": "AI_FIRE", "service_room": "AI_MEP",
               "staircase": "AI_STAIR", "lift_shaft": "AI_LIFT", "condenser": "AI_SHAFT",
               "double_cutout": "AI_BALCONY", "triple_cutout": "AI_UNKNOWN", "void": "AI_VOID_AREA",
               "shaft_room": "AI_MEP", "wall_cutout": "AI_CUTOUT", "cutout_above": "AI_SHAFT",
               "circulation": "AI_CIRC"}
    out = []
    for key, group, header in CATEGORIES:
        if key in MAP_SKIP or group == "derived":
            continue
        out.append((key, header.replace("\n", " "), GROUP_NAME[group], typical.get(key, "AI_SHAFT"),
                    "Y" if group == "opening" else "N", "Y" if group == "vertical" else "N"))
    out.append(("usable", "USABLE AREA", "Usable", "(any)", "N", "N"))
    return out


def _statement_keys(data):
    """The category columns some plate uses, in CATEGORIES order, with the
    fixed head (gross) and tail (the derived figures)."""
    used = set()
    for t in data["towers"]:
        for r in t["rows"]:
            used |= {k for k, v in (r.get("canon") or {}).items() if isinstance(v, (int, float)) and abs(v) > 0.004}
    keys = [k for k, g, _ in CATEGORIES if k not in MAP_SKIP and g not in ("area", "derived") and k in used]
    return ["gross"] + keys + list(DERIVED)


def write_area_statement(area_data: Dict, out_path: str, drawing_name: str = "",
                         project_name: str = "", cfg=None, template=None) -> Dict:
    """The Area Statement workbook: one sheet, every figure a formula over
    the AREA REGISTER at the foot of it (the measured elements, in yellow),
    through defined names - the statement, a quick lookup, the register with
    an outline (plate > group > column > element) and a by-layer table, and
    the working data (settings, floor list, category map, group titles).

    INDEX/MATCH and SUMIFS only, so it opens in any Excel. build_statement
    decides what goes where; this writes it."""
    from openpyxl.workbook.defined_name import DefinedName
    data = build_statement(area_data, drawing_name, cfg, template)
    sf, sqft = data["sale_factor"], data["sqft"]
    keys = _statement_keys(data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Area Statement"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False
    sheet_ref = "'Area Statement'!"
    names = {}

    def name(nm, ref):
        names[nm] = sheet_ref + ref

    def col(i):
        return get_column_letter(FIRST_COL + i)
    last_col = col(len(keys) + 1)                     # sl, floor, then the keys
    key_col = {k: col(2 + i) for i, k in enumerate(keys)}

    def put(ref, value, font=None, fill=None, align=None, fmt=None, border=True):
        c = ws[ref]
        if type(c).__name__ != "MergedCell":
            c.value = value
        if font is not None:
            c.font = font
        if fill is not None:
            c.fill = fill
        if align is not None:
            c.alignment = align
        if fmt is not None:
            c.number_format = fmt
        if border:
            c.border = BORDER
        return c

    # -- 1. title ------------------------------------------------------------
    put("B2", '="Project Name: "&Project', Font(bold=True, size=16, color=NAVY), align=LEFT, border=False)
    put("B3", '="Area Statement  -  "&Drawing&"  -  all areas in sq.m unless stated  -  sale area = plinth / "'
              '&TEXT(SaleFactor,"0.00")', Font(italic=True, size=10, color="465A7A"), border=False)

    # -- 2. the statement, one block per tower -------------------------------
    floors_all = []                                  # (floor label, plate, order, usage, refuge)
    order = 0
    row = 5
    summary_rows = []
    for t in data["towers"]:
        offset = order
        ws.merge_cells("B{0}:{1}{0}".format(row, last_col))
        put("B{0}".format(row), "=Tower" if len(data["towers"]) == 1 else t["title"],
            Font(bold=True, size=14, color="FFFFFF"), HEAD_FILL, border=False)
        ws.row_dimensions[row].height = 24
        row += 2
        hrow, krow = row, row + 2                     # header rows (two), then the key row
        # group titles over the columns
        put("B{0}".format(hrow), "Sl No", WHITE_BOLD, HEAD_FILL, CENTER)
        put("C{0}".format(hrow), "FLOORS", WHITE_BOLD, HEAD_FILL, CENTER)
        ws.merge_cells("B{0}:B{1}".format(hrow, hrow + 1))
        ws.merge_cells("C{0}:C{1}".format(hrow, hrow + 1))
        put("D{0}".format(hrow), '=INDEX(M_Header,MATCH(D${0},M_Key,0))'.format(krow), WHITE_BOLD, HEAD_FILL, CENTER)
        ws.merge_cells("D{0}:D{1}".format(hrow, hrow + 1))
        groups = []
        for k in keys[1:]:
            g = GROUP_OF.get(k)
            if groups and groups[-1][0] == g:
                groups[-1][2] = key_col[k]
            else:
                groups.append([g, key_col[k], key_col[k]])
        for g, c0, c1 in groups:
            if c1 != c0:
                ws.merge_cells("{0}{2}:{1}{2}".format(c0, c1, hrow))
            if g == "derived":
                f = '=INDEX(G_Title,MATCH("Derived",G_Group,0))'
            else:
                f = '=INDEX(G_Title,MATCH(INDEX(M_Group,MATCH({0}${1},M_Key,0)),G_Group,0))'.format(c0, krow)
            put("{0}{1}".format(c0, hrow), f, Font(bold=True, size=10, color=NAVY), GROUP_FILL, CENTER)
        for k in keys[1:]:
            c = key_col[k]
            if GROUP_OF.get(k) == "derived":
                put("{0}{1}".format(c, hrow + 1), HEADER_OF[k].replace("\n", " "), WHITE_BOLD, HEAD_FILL, CENTER)
            else:
                put("{0}{1}".format(c, hrow + 1), '=INDEX(M_Header,MATCH({0}${1},M_Key,0))'.format(c, krow),
                    WHITE_BOLD, HEAD_FILL, CENTER)
        for i in range(len(keys) + 2):
            for r in (hrow, hrow + 1):
                ws.cell(row=r, column=FIRST_COL + i).border = BORDER
        ws.row_dimensions[hrow].height = 20
        ws.row_dimensions[hrow + 1].height = 48
        for k in keys:
            put("{0}{1}".format(key_col[k], krow), k, KEY_FONT, align=CENTER, border=False)
        ws.row_dimensions[krow].height = 11
        first = krow + 1
        r = first
        for rec in t["rows"]:
            order += 1
            floors_all.append((rec["level"], rec.get("plate") or rec["level"], order,
                               "Refuge" if "(Refuge)" in rec["level"] else "Office", "(Refuge)" in rec["level"]))
            refuge = "(Refuge)" in rec["level"]
            put("B{0}".format(r), "=ROW()-ROW($B${0})".format(krow), align=CENTER, fmt="0")
            sl = "B{0}".format(r) if not offset else "(B{0}+{1})".format(r, offset)
            put("C{0}".format(r), "=INDEX(F_Floor,MATCH({0},F_Order,0))".format(sl),
                Font(bold=True, color=REFUGE_RED) if refuge else None, fmt="#,##0.00")
            plate = "INDEX(F_Plate,MATCH($C{0},F_Floor,0))".format(r)
            for k in keys:
                c = key_col[k]
                ref = "{0}{1}".format(c, r)
                if k == "total_shafts":
                    f = "=SUMPRODUCT(SUMIFS(R_Area,R_Plate,{0},R_Key,M_Key)*(M_Deducted=\"Y\"))".format(plate)
                elif k == "net_bua":
                    f = "={0}{1}-{2}{1}".format(key_col["gross"], r, key_col["total_shafts"])
                elif k == "plinth":
                    f = "={0}{1}-SUMPRODUCT(SUMIFS(R_Area,R_Plate,{2},R_Key,M_Key)*(M_PlinthLess=\"Y\"))".format(
                        key_col["net_bua"], r, plate)
                elif k == "sale_sqm":
                    f = "={0}{1}/SaleFactor".format(key_col["plinth"], r)
                elif k == "sale_sft":
                    f = "={0}{1}*SqftPerSqm".format(key_col["sale_sqm"], r)
                elif k == "eff_sa":
                    f = "=IFERROR({0}{1}/{2}{1},0)".format(key_col["gross"], r, key_col["sale_sqm"])
                elif k == "eff_pa":
                    f = "=IFERROR({0}{1}/{2}{1},0)".format(key_col["gross"], r, key_col["plinth"])
                else:
                    f = "=SUMIFS(R_Area,R_Plate,{0},R_Key,{1}${2})".format(plate, c, krow)
                cell = put(ref, f, fmt="0.00" if k in ("eff_sa", "eff_pa") else "#,##0.00")
                if k in KEY_FILLS:
                    cell.fill = KEY_FILLS[k]
                elif GROUP_OF.get(k) == "info":
                    cell.fill = INFO_FILL
                if refuge:
                    cell.font = Font(color=REFUGE_RED)
            r += 1
        last = r - 1
        put("C{0}".format(r), '="TOTAL ("&COUNT(B{0}:B{1})&" floors)"'.format(first, last),
            Font(bold=True, color=NAVY), TOTAL_BAND)
        put("B{0}".format(r), None, fill=TOTAL_BAND)
        for k in keys:
            c = key_col[k]
            if k == "eff_sa":
                f = "=IFERROR({0}{1}/{2}{1},0)".format(key_col["gross"], r, key_col["sale_sqm"])
            elif k == "eff_pa":
                f = "=IFERROR({0}{1}/{2}{1},0)".format(key_col["gross"], r, key_col["plinth"])
            else:
                f = "=SUM({0}{1}:{0}{2})".format(c, first, last)
            put("{0}{1}".format(c, r), f, Font(bold=True, color=NAVY), TOTAL_BAND,
                fmt="0.00" if k in ("eff_sa", "eff_pa") else "#,##0.00")
        total_row = r
        r += 1
        put("C{0}".format(r), "GROSS AREA (sq.ft)", Font(bold=True, color="1F4E79"), border=False)
        put("{0}{1}".format(key_col["gross"], r), "={0}{1}*SqftPerSqm".format(key_col["gross"], total_row),
            Font(bold=True, color="1F4E79"), fmt="#,##0.00", border=False)
        summary_rows.append((t["title"], total_row))
        row = r + 2
    # the summary
    ws.merge_cells("B{0}:F{0}".format(row))
    put("B{0}".format(row), "SUMMARY", Font(bold=True, size=12, color=NAVY), GROUP_FILL, border=False)
    row += 1
    for i, h in enumerate(["S.No", "Tower", "GBUA (sq.m)", "GBUA (sq.ft)", "Sale area (sq.ft)"]):
        put("{0}{1}".format(col(i), row), h, WHITE_BOLD, HEAD_FILL, CENTER)
    row += 1
    sfirst = row
    for i, (title, trow) in enumerate(summary_rows, start=1):
        put("B{0}".format(row), i)
        put("C{0}".format(row), "=Tower" if len(summary_rows) == 1 else title)
        put("D{0}".format(row), "={0}{1}".format(key_col["gross"], trow), fmt="#,##0.00")
        put("E{0}".format(row), "={0}{1}*SqftPerSqm".format(key_col["gross"], trow), fmt="#,##0.00")
        put("F{0}".format(row), "={0}{1}".format(key_col["sale_sft"], trow), fmt="#,##0.00")
        row += 1
    put("B{0}".format(row), None, fill=TOTAL_BAND)
    put("C{0}".format(row), "Total", Font(bold=True, color=NAVY), TOTAL_BAND)
    for c in "DEF":
        put("{0}{1}".format(c, row), "=SUM({0}{1}:{0}{2})".format(c, sfirst, row - 1),
            Font(bold=True, color=NAVY), TOTAL_BAND, fmt="#,##0.00")
    ws.print_area = "B2:{0}{1}".format(last_col, row)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    row += 3

    # -- 3. quick lookup -----------------------------------------------------
    ws.merge_cells("B{0}:I{0}".format(row))
    put("B{0}".format(row), "QUICK LOOKUP    type in the yellow cells; Floor blank = whole tower",
        Font(bold=True, size=12, color="FFFFFF"), LOOKUP_FILL, border=False)
    ws.row_dimensions[row].height = 22
    row += 1
    lk = row
    first_floor = floors_all[0][0] if floors_all else ""
    first_key = next((k for k in keys[1:] if GROUP_OF.get(k) == "opening"), keys[1] if len(keys) > 1 else "gross")
    for i, (label, value, nm) in enumerate([("Floor", first_floor, "LkFloor"),
                                            ("Element name (or part of it)", "TOILET", "LkElement"),
                                            ("Category key (see map)", first_key, "LkCategory")]):
        put("C{0}".format(lk + i), label, Font(bold=True))
        put("D{0}".format(lk + i), value, fill=INPUT_FILL)
        name(nm, "$D${0}".format(lk + i))
    plate_of = 'INDEX(F_Plate,MATCH(LkFloor,F_Floor,0))'
    look = [("Element area (sq.m)",
             '=IF(LkFloor="",SUMPRODUCT(SUMIFS(R_Area,R_Plate,F_Plate,R_Key,"<>",R_Name,"*"&LkElement&"*")),'
             'SUMIFS(R_Area,R_Plate,{0},R_Key,"<>",R_Name,"*"&LkElement&"*"))'.format(plate_of), "#,##0.00"),
            ("Element nos",
             '=IF(LkFloor="",SUMPRODUCT(SUMIFS(R_Nos,R_Plate,F_Plate,R_Key,"<>",R_Name,"*"&LkElement&"*")),'
             'SUMIFS(R_Nos,R_Plate,{0},R_Key,"<>",R_Name,"*"&LkElement&"*"))'.format(plate_of), "0"),
            ("Category area (sq.m)",
             '=IF(LkFloor="",SUMPRODUCT(SUMIFS(R_Area,R_Plate,F_Plate,R_Key,LkCategory)),'
             'SUMIFS(R_Area,R_Plate,{0},R_Key,LkCategory))'.format(plate_of), "#,##0.00"),
            ("Area (sq.ft) of the element", "=H{0}*SqftPerSqm".format(lk), "#,##0.00")]
    for i, (label, f, fmt) in enumerate(look):
        put("F{0}".format(lk + i), label)
        put("G{0}".format(lk + i), None)
        put("H{0}".format(lk + i), f, Font(bold=True), fmt=fmt)
        put("I{0}".format(lk + i), None)
        ws.merge_cells("F{0}:G{0}".format(lk + i))
        ws.merge_cells("H{0}:I{0}".format(lk + i))
    row = lk + 6

    # -- 4. the area register ------------------------------------------------
    ws.merge_cells("B{0}:I{0}".format(row))
    put("B{0}".format(row), "AREA REGISTER    grouped by statement column - buttons 1 / 2 / 3 on the left: "
                            "groups / columns / every element", Font(bold=True, size=12, color=NAVY), GROUP_FILL,
        border=False)
    put("K{0}".format(row), "BY LAYER", Font(bold=True, size=12, color=NAVY), GROUP_FILL, border=False)
    ws.merge_cells("K{0}:M{0}".format(row))
    ws.row_dimensions[row].height = 22
    row += 1
    put("C{0}".format(row), "Yellow = measured (edit Nos, Area or Key to correct). Key decides the statement column.",
        Font(italic=True, size=9, color="465A7A"), border=False)
    put("K{0}".format(row), "Plate (blank = all)", Font(bold=True, size=9), border=False)
    put("L{0}".format(row), "", fill=INPUT_FILL)
    name("LayPlate", "$L${0}".format(row))
    row += 1
    for i, h in enumerate(["Element", "Layer", "Nos", "Area (sq.m)", "Area (sq.ft)", "Plate", "Key"]):
        put("{0}{1}".format(get_column_letter(3 + i), row), h, WHITE_BOLD, HEAD_FILL, CENTER)
    for i, h in enumerate(["Layer", "Nos", "Area (sq.m)"]):
        put("{0}{1}".format(get_column_letter(11 + i), row), h, WHITE_BOLD, HEAD_FILL, CENTER)
    row += 1
    reg_first = row
    layer_row = row
    # the plates, each once, with their elements grouped
    plates = []
    seen_plates = set()
    for t in data["towers"]:
        for rec in t["rows"]:
            pl = rec.get("plate") or rec["level"]
            if pl not in seen_plates:
                seen_plates.add(pl)
                plates.append((pl, rec["gross"]))
    rooms_by_plate = {}
    for rm in data["rooms"]:
        rooms_by_plate.setdefault(rm.get("plate") or rm.get("level"), []).append(rm)
    many = len(plates) > 1
    base_level = 1 if many else 0
    layer_totals = {}

    def element_rows(items):
        """(name, layer, nos, area) per (name, layer), largest first."""
        agg = {}
        for rm in items:
            lay = _family_layer_of(rm.get("family") or "", rm.get("layer") or "")
            k = (rm.get("name") or "", lay)
            a = agg.setdefault(k, [0, 0.0])
            a[0] += 1
            a[1] += float(rm.get("area_sqm") or 0.0)
        out = [(k[0], k[1], v[0], round(v[1], 2)) for k, v in agg.items()]
        out.sort(key=lambda e: -e[3])
        return out

    def write_subtotal(r, label, level, fill, font):
        put("C{0}".format(r), label, font, fill)
        put("D{0}".format(r), None, fill=fill)
        put("E{0}".format(r), None, font, fill, fmt="0")
        put("F{0}".format(r), None, font, fill, fmt="#,##0.00")
        put("G{0}".format(r), "=F{0}*SqftPerSqm".format(r), font, fill, fmt="#,##0.00")
        put("H{0}".format(r), None, fill=fill)
        put("I{0}".format(r), None, fill=fill)
        ws.row_dimensions[r].outlineLevel = level

    def write_element(r, plate, key, name_, layer, nos, area, level):
        put("C{0}".format(r), name_)
        put("D{0}".format(r), layer)
        put("E{0}".format(r), nos, fill=INPUT_FILL, fmt="0")
        put("F{0}".format(r), area, fill=INPUT_FILL, fmt="#,##0.00")
        put("G{0}".format(r), "=F{0}*SqftPerSqm".format(r), fmt="#,##0.00")
        put("H{0}".format(r), plate)
        put("I{0}".format(r), key, REG_KEY_FONT, INPUT_FILL)
        ws.row_dimensions[r].outlineLevel = level
        layer_totals[layer] = layer_totals.get(layer, 0.0) + area

    for pl, gross in plates:
        if many:
            put("C{0}".format(row), "PLATE:  " + pl, Font(bold=True, size=12, color=NAVY), PLATE_FILL)
            put("D{0}".format(row), '="serves "&COUNTIF(F_Plate,H{0})&" floor"&IF(COUNTIF(F_Plate,H{0})>1,"s","")'.format(row),
                Font(italic=True, size=10, color=NAVY), PLATE_FILL)
            put("E{0}".format(row), None, fill=PLATE_FILL)
            put("F{0}".format(row), '=SUMIFS(R_Area,R_Plate,H{0},R_Key,"gross")'.format(row),
                Font(bold=True, size=12, color=NAVY), PLATE_FILL, fmt="#,##0.00")
            put("G{0}".format(row), None, fill=PLATE_FILL)
            put("H{0}".format(row), pl, Font(bold=True, size=12, color=NAVY), PLATE_FILL)
            put("I{0}".format(row), None, fill=PLATE_FILL)
            ws.row_dimensions[row].height = 20
            ws.row_dimensions[row].outlineLevel = 0
            row += 1
        items = rooms_by_plate.get(pl) or []
        for gkey, gtitle in GROUP_KEYS:
            if gkey == "derived":
                continue
            if gkey == "area":
                members = [("gross", [("FLOOR PLATE (Pline)", "AI_FLOOR", 1, round(float(gross), 2))])]
            elif gkey == "usable":
                els = element_rows([rm for rm in items if not rm.get("key")])
                members = [("usable", els)] if els else []
            else:
                members = []
                for k, g, _h in CATEGORIES:
                    if g != gkey:
                        continue
                    els = element_rows([rm for rm in items if rm.get("key") == k])
                    if els:
                        members.append((k, els))
            if not members:
                continue
            grow = row
            row += 1
            for k, els in members:
                if gkey in ("area", "usable"):
                    for name_, lay, nos, area in els:
                        write_element(row, pl, k, name_, lay, nos, area, base_level + 1)
                        row += 1
                    continue
                crow = row
                row += 1
                for name_, lay, nos, area in els:
                    write_element(row, pl, k, name_, lay, nos, area, base_level + 2)
                    row += 1
                write_subtotal(crow, HEADER_OF[k].replace("\n", " "), base_level + 1, COLUMN_FILL, Font(bold=True))
                ws["E{0}".format(crow)] = "=SUBTOTAL(9,E{0}:E{1})".format(crow + 1, row - 1)
                ws["F{0}".format(crow)] = "=SUBTOTAL(9,F{0}:F{1})".format(crow + 1, row - 1)
            write_subtotal(grow, gtitle, base_level, GROUP_FILL, Font(bold=True, size=11, color=NAVY))
            ws["E{0}".format(grow)] = "=SUBTOTAL(9,E{0}:E{1})".format(grow + 1, row - 1)
            ws["F{0}".format(grow)] = "=SUBTOTAL(9,F{0}:F{1})".format(grow + 1, row - 1)
    reg_last = max(row - 1, reg_first)
    for nm, c in (("R_Name", "C"), ("R_Layer", "D"), ("R_Nos", "E"), ("R_Area", "F"), ("R_Plate", "H"), ("R_Key", "I")):
        name(nm, "${0}${1}:${0}${2}".format(c, reg_first, reg_last))
    # by layer, beside the register
    lr = layer_row
    for lay, _a in sorted(layer_totals.items(), key=lambda kv: -kv[1]):
        put("K{0}".format(lr), lay)
        put("L{0}".format(lr), '=SUMIFS(R_Nos,R_Layer,K{0},R_Plate,IF(LayPlate="","*",LayPlate))'.format(lr), fmt="0")
        put("M{0}".format(lr), '=SUMIFS(R_Area,R_Layer,K{0},R_Plate,IF(LayPlate="","*",LayPlate))'.format(lr),
            fmt="#,##0.00")
        lr += 1
    row = max(row, lr) + 2

    # -- 5. working data (grouped, collapsed) --------------------------------
    ws.merge_cells("B{0}:{1}{0}".format(row, last_col))
    put("B{0}".format(row), "WORKING DATA    settings, floor list, category map - click [+] on the left to open",
        Font(bold=True, size=12, color="FFFFFF"), LOOKUP_FILL, border=False)
    ws.row_dimensions[row].height = 22
    row += 1
    wd_first = row
    put("C{0}".format(row), "Settings", Font(bold=True, color=NAVY), border=False)
    row += 1
    tower_title = data["towers"][0]["title"] if data["towers"] else "TOWER 1"
    for label, value, nm, fmt in (("Project name", (project_name or "").strip() or "-", "Project", None),
                                  ("Drawing", drawing_name or "the drawing", "Drawing", None),
                                  ("Tower", tower_title, "Tower", None),
                                  ("Sale area factor", float(sf), "SaleFactor", "0.00"),
                                  ("sq.ft per sq.m", float(sqft), "SqftPerSqm", "0.000")):
        put("C{0}".format(row), label, Font(bold=True))
        put("D{0}".format(row), value, fill=INPUT_FILL, fmt=fmt)
        name(nm, "$D${0}".format(row))
        row += 1
    put("C{0}".format(row), "Floor list  (each floor -> the plate in the register; a typical plate serves "
                            "several floors)", Font(bold=True, color=NAVY), border=False)
    row += 1
    for i, h in enumerate(["Floor", "Plate", "Order", "Usage"]):
        put("{0}{1}".format(get_column_letter(3 + i), row), h, WHITE_BOLD, HEAD_FILL, CENTER)
    row += 1
    f_first = row
    for label, pl, order_, usage, refuge in floors_all:
        put("C{0}".format(row), label, Font(color=REFUGE_RED) if refuge else None, INPUT_FILL)
        put("D{0}".format(row), pl, fill=INPUT_FILL)
        put("E{0}".format(row), order_, fill=INPUT_FILL, fmt="0")
        put("F{0}".format(row), usage, fill=INPUT_FILL)
        row += 1
    for _ in range(FLOOR_SPARE_ROWS):
        for c in "CDEF":
            put("{0}{1}".format(c, row), None, fill=INPUT_FILL)
        row += 1
    f_last = row - 1
    name("F_Floor", "$C${0}:$C${1}".format(f_first, f_last))
    name("F_Plate", "$D${0}:$D${1}".format(f_first, f_last))
    name("F_Order", "$E${0}:$E${1}".format(f_first, f_last))
    row += 1
    put("C{0}".format(row), "Category map  (Header renames a statement column; Deducted Y = in Total Shafts; "
                            "PlinthLess Y = off the plinth)", Font(bold=True, color=NAVY), border=False)
    row += 1
    for i, h in enumerate(["Key", "Header", "Group", "Typical layer", "Deducted", "PlinthLess"]):
        put("{0}{1}".format(get_column_letter(3 + i), row), h, WHITE_BOLD, HEAD_FILL, CENTER)
    row += 1
    m_first = row
    for key, header, group, layer, ded, pless in _map_rows():
        for i, v in enumerate((key, header, group, layer, ded, pless)):
            put("{0}{1}".format(get_column_letter(3 + i), row), v, fill=INPUT_FILL)
        row += 1
    m_last = row - 1
    name("M_Key", "$C${0}:$C${1}".format(m_first, m_last))
    name("M_Header", "$D${0}:$D${1}".format(m_first, m_last))
    name("M_Group", "$E${0}:$E${1}".format(m_first, m_last))
    name("M_Deducted", "$G${0}:$G${1}".format(m_first, m_last))
    name("M_PlinthLess", "$H${0}:$H${1}".format(m_first, m_last))
    row += 1
    put("C{0}".format(row), "Group", WHITE_BOLD, HEAD_FILL, CENTER)
    put("D{0}".format(row), "Group title", WHITE_BOLD, HEAD_FILL, CENTER)
    row += 1
    g_first = row
    for gkey, gtitle in GROUP_KEYS:
        put("C{0}".format(row), GROUP_NAME[gkey], fill=INPUT_FILL)
        put("D{0}".format(row), gtitle, fill=INPUT_FILL)
        row += 1
    g_last = row - 1
    name("G_Group", "$C${0}:$C${1}".format(g_first, g_last))
    name("G_Title", "$D${0}:$D${1}".format(g_first, g_last))
    for r in range(wd_first, row):
        ws.row_dimensions[r].outlineLevel = 1
        ws.row_dimensions[r].hidden = True

    # names, widths
    for nm, ref in names.items():
        dn = DefinedName(nm, attr_text=ref)
        try:
            wb.defined_names[nm] = dn
        except TypeError:                             # older openpyxl
            wb.defined_names.append(dn)
    ws.column_dimensions["A"].width = 2.0
    ws.column_dimensions["B"].width = 7.0
    ws.column_dimensions["C"].width = 34.0
    ws.column_dimensions["D"].width = 17.0
    for i in range(4, len(keys) + 2):
        ws.column_dimensions[col(i)].width = 12.0
    ws.column_dimensions["K"].width = max(ws.column_dimensions["K"].width or 0, 18.0)
    ws.column_dimensions["L"].width = max(ws.column_dimensions["L"].width or 0, 17.0)
    wb.save(out_path)
    data["columns"] = [HEADER_OF[k] for k in keys]
    data["statement_keys"] = keys
    return data


# -- the statement for the browser -------------------------------------------
def statement_view(area_data: Dict, drawing_name: str = "", cfg=None, template=None) -> Dict:
    """The Area Statement as data, for the page: towers, rows, columns and
    totals, with the workbook's formulas worked out."""
    data = build_statement(area_data, drawing_name, cfg, template)
    sf, sqft = data["sale_factor"], data["sqft"]
    sh = _Sheet(_active_columns(data), data["template"])
    tpl = data["template"]
    towers, summary = [], []
    for t in data["towers"]:
        rows = []
        total = {k: 0.0 for k in sh.keys if k not in ("sl", "floor")}
        for sl, rec in enumerate(t["rows"], start=1):
            cells = dict(rec["values"])
            cells["gross"] = rec["gross"]
            cells.update(derive(rec["values"], rec["gross"], tpl, sf, sqft))
            row = {"sl": sl, "floor": rec["level"], "plate": rec.get("plate", ""),
                   "refuge": "(Refuge)" in rec["level"], "canon": rec.get("canon", {})}
            for k in sh.keys:
                if k in ("sl", "floor"):
                    continue
                row[k] = round(cells.get(k, 0.0), 2)
                if k not in ("eff_sa", "eff_pa"):
                    total[k] += cells.get(k, 0.0)
            rows.append(row)
        total = {k: round(v, 2) for k, v in total.items()}
        total["eff_sa"] = round(total["gross"] / total["sale_sqm"], 2) if total.get("sale_sqm") else 0.0
        total["eff_pa"] = round(total["gross"] / total["plinth"], 2) if total.get("plinth") else 0.0
        total["floor"] = "TOTAL ({0} floors)".format(len(rows))
        towers.append({"title": t["title"], "rows": rows, "total": total,
                       "gross_sft": round(total["gross"] * sqft, 2)})
        summary.append({"tower": t["title"], "gbua_sqm": total["gross"],
                        "gbua_sft": round(total["gross"] * sqft, 2), "sale_sft": total["sale_sft"]})
    return {
        "drawing": drawing_name,
        "template": tpl["name"],
        "columns": [{"key": c["key"], "label": c["header"],
                     "group": GROUP_TITLE.get(c.get("group")) if c.get("group") else None}
                    for c in sh.cols],
        "towers": towers,
        "summary": summary,
        "summary_total": {"gbua_sqm": round(sum(s["gbua_sqm"] for s in summary), 2),
                          "gbua_sft": round(sum(s["gbua_sft"] for s in summary), 2),
                          "sale_sft": round(sum(s["sale_sft"] for s in summary), 2)},
        "rooms": data["rooms"],
        "sale_factor": sf, "sqft": sqft,
    }
